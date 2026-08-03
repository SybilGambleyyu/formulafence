"""Semantic workbook diffing, downstream impact, and intrinsic-risk detection."""

from __future__ import annotations

from collections import defaultdict, deque
from collections.abc import Iterable, Mapping
from dataclasses import dataclass

from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter

from formulafence.formulas import resolve_3d_reference
from formulafence.models import (
    AlignmentSnapshot,
    BorderSnapshot,
    CellHyperlinkSnapshot,
    CellKey,
    CellSnapshot,
    Change,
    ChartDefinitionSnapshot,
    ConditionalFormattingExtensionSnapshot,
    ConditionalFormattingSnapshot,
    CustomDataStoreSnapshot,
    CustomWorkbookViewSnapshot,
    DataValidationSnapshot,
    DiffReport,
    DigitalSignatureSnapshot,
    ExternalDataConnectionSnapshot,
    ExternalLinkPackageSnapshot,
    ExternalRelationshipSnapshot,
    ExternalWorkbookLinkSurfaceSnapshot,
    FillSnapshot,
    FilterVisibilitySnapshot,
    Finding,
    FontSnapshot,
    FormulaCachedResultSnapshot,
    FormulaDdeLinkSnapshot,
    FormulaDefinedXlmActionSnapshot,
    FormulaDefinedXlmEnvironmentInformationSnapshot,
    FormulaDefinedXlmEvaluationSnapshot,
    FormulaDefinedXlmGetCellSnapshot,
    FormulaDefinedXlmRegistrationSnapshot,
    FormulaEnvironmentInformationSnapshot,
    FormulaExternalActionSnapshot,
    FormulaFenceError,
    IgnoredErrorSnapshot,
    LegacyCommentSnapshot,
    NamedSheetViewSnapshot,
    NumberFormatSnapshot,
    OfficeCustomFunctionSnapshot,
    OfficeWebAddinSnapshot,
    PivotCacheRefreshSnapshot,
    PivotTableDefinitionSnapshot,
    PowerPivotDataModelSnapshot,
    PowerQuerySnapshot,
    ProtectionCredentialSnapshot,
    ProtectionOpaqueMetadataSnapshot,
    PythonInExcelSnapshot,
    QueryTableRefreshSnapshot,
    RibbonCustomizationSnapshot,
    RichDataSnapshot,
    RichTextRunEntry,
    RichTextRunSnapshot,
    ScenarioManagerSnapshot,
    SharedWorkbookRevisionSnapshot,
    SlicerTimelineCacheSnapshot,
    TableStyleControlsSnapshot,
    ThreadedCommentSnapshot,
    UnqualifiedRuntimeFunctionSnapshot,
    WhatIfDataTableSnapshot,
    WorkbookSnapshot,
    WorkbookThemeSnapshot,
    WorksheetCodeResourceRegistrationSnapshot,
    WorksheetDimensionSnapshot,
    WorksheetDisplaySnapshot,
    WorksheetDrawingShapeSnapshot,
    WorksheetEmbeddedControlSnapshot,
    WorksheetImageSnapshot,
    WorksheetPrintLayoutSnapshot,
    WorksheetSparklineSnapshot,
    XlmAutomaticMacroBindingSnapshot,
    XlmMacroSheetSnapshot,
    XmlMappingSnapshot,
)

_IMPACT_SAMPLE_SIZE = 20
DEFAULT_MAX_CHANGE_ANALYSIS_STATES = 100_000
_IMPACT_NODE_LIMIT = DEFAULT_MAX_CHANGE_ANALYSIS_STATES


@dataclass
class ChangeAnalysisBudget:
    """One bounded pool for local dependency evidence in a comparison.

    Each source cell and each reachable dependent that FormulaFence retains
    while tracing a change consumes one state.  The budget is deliberately
    shared by every local-impact query in a comparison: a wide edit set must
    not multiply the existing per-source traversal bound into impractical CI
    work or report memory.
    """

    max_states: int
    scope: str
    error_type: type[FormulaFenceError] = FormulaFenceError
    state_count: int = 0

    def consume(self) -> None:
        """Record a source-to-reachable state or stop before overage work."""
        if self.state_count >= self.max_states:
            raise self.error_type(
                f"{self.scope} change analysis exceeds "
                f"max_change_analysis_states={self.max_states}."
            )
        self.state_count += 1


@dataclass(frozen=True)
class ImpactAnalysis:
    """A deterministic, shortest-path view of explicit downstream dependencies."""

    impacted: frozenset[CellKey]
    paths: Mapping[CellKey, tuple[CellKey, ...]]
    truncated: bool


@dataclass(frozen=True)
class _ImpactPaths(Mapping[CellKey, tuple[CellKey, ...]]):
    """Reconstruct only the impact paths that a caller actually renders.

    A long formula chain can have many reachable nodes and therefore many
    quadratic-length path prefixes. Review output deliberately serializes only
    a small deterministic sample, so eagerly constructing every prefix would
    turn a bounded graph walk into an avoidable allocation spike.
    """

    parents: Mapping[CellKey, CellKey | None]
    targets: frozenset[CellKey]

    def __getitem__(self, target: CellKey) -> tuple[CellKey, ...]:
        if target not in self.targets:
            raise KeyError(target)
        path: list[CellKey] = []
        current: CellKey | None = target
        while current is not None:
            path.append(current)
            current = self.parents[current]
        return tuple(reversed(path))

    def __iter__(self):
        return iter(self.targets)

    def __len__(self) -> int:
        return len(self.targets)


def _cell_change_kind(
    before: CellSnapshot | None, after: CellSnapshot | None
) -> tuple[str, str] | None:
    if before is None and after is None:  # pragma: no cover - impossible by caller contract
        return None
    if before is None:
        return ("formula_added", "low") if after and after.is_formula else ("value_added", "low")
    if after is None:
        return ("formula_removed", "high") if before.is_formula else ("value_removed", "low")
    if before.is_formula and not after.is_formula:
        return "formula_to_value", "high"
    if not before.is_formula and after.is_formula:
        return "value_to_formula", "medium"
    if before.is_formula and after.is_formula and before.formula != after.formula:
        return "formula_changed", "medium"
    if before.cell_type != after.cell_type or before.value != after.value:
        return "value_changed", "low"
    return None


def analyze_downstream_impact(
    location: CellKey,
    *snapshots: WorkbookSnapshot,
    node_limit: int = _IMPACT_NODE_LIMIT,
    state_budget: ChangeAnalysisBudget | None = None,
) -> ImpactAnalysis:
    """Find formula cells affected through explicit dependency paths.

    Range references are checked lazily instead of expanded, so a formula such as
    `SUM(A:A)` does not turn into a million in-memory graph edges.
    """
    for snapshot in snapshots:
        snapshot.require_full_inspection("Downstream impact analysis")
    if state_budget is not None:
        state_budget.consume()
    queue: deque[CellKey] = deque([location])
    visited: set[CellKey] = {location}
    impacted: set[CellKey] = set()
    parents: dict[CellKey, CellKey | None] = {location: None}
    truncated = False
    while queue:
        current = queue.popleft()
        dependents: set[CellKey] = set()
        for snapshot in snapshots:
            dependents.update(snapshot.direct_dependents(current))
        for dependent in sorted(dependents, key=_location_sort_key):
            if dependent in visited:
                continue
            if state_budget is not None:
                state_budget.consume()
            visited.add(dependent)
            impacted.add(dependent)
            parents[dependent] = current
            if len(visited) >= node_limit:
                truncated = True
                queue.clear()
                break
            queue.append(dependent)
    impacted_nodes = frozenset(impacted)
    return ImpactAnalysis(
        impacted_nodes,
        _ImpactPaths(parents, impacted_nodes),
        truncated,
    )


def downstream_impact(
    location: CellKey,
    *snapshots: WorkbookSnapshot,
    node_limit: int = _IMPACT_NODE_LIMIT,
) -> tuple[set[CellKey], bool]:
    """Backward-compatible shortcut for callers that need counts only."""
    analysis = analyze_downstream_impact(location, *snapshots, node_limit=node_limit)
    return set(analysis.impacted), analysis.truncated


def _serialise_impact_paths(
    targets: Iterable[CellKey],
    paths: Mapping[CellKey, tuple[CellKey, ...]],
) -> list[dict[str, object]]:
    from formulafence.models import display_location

    return [
        {
            "target": display_location(target),
            "path": [display_location(step) for step in paths[target]],
        }
        for target in targets
    ]


def _location_sort_key(location: CellKey | None) -> tuple[str, int, int]:
    if location is None:
        return ("", 0, 0)
    sheet, coordinate = location
    row, column = coordinate_to_tuple(coordinate)
    return sheet.casefold(), row, column


def _adjacent(location: CellKey, row_delta: int, column_delta: int) -> CellKey | None:
    sheet, coordinate = location
    row, column = coordinate_to_tuple(coordinate)
    target_row = row + row_delta
    target_column = column + column_delta
    if target_row < 1 or target_column < 1:
        return None
    return sheet, f"{get_column_letter(target_column)}{target_row}"


def _formula_pattern_findings(
    snapshot: WorkbookSnapshot, changed_locations: set[CellKey]
) -> list[Finding]:
    """Detect a changed formula that breaks a stable immediate peer pattern."""
    findings: list[Finding] = []
    for location in sorted(changed_locations, key=_location_sort_key):
        cell = snapshot.cells.get(location)
        if cell is None or not cell.is_formula or cell.formula_fingerprint is None:
            continue
        for orientation, before_delta, after_delta in (
            ("column", (-1, 0), (1, 0)),
            ("row", (0, -1), (0, 1)),
        ):
            previous_location = _adjacent(location, *before_delta)
            next_location = _adjacent(location, *after_delta)
            if previous_location is None or next_location is None:
                continue
            previous = snapshot.cells.get(previous_location)
            following = snapshot.cells.get(next_location)
            if (
                previous is not None
                and following is not None
                and previous.is_formula
                and following.is_formula
                and previous.formula_fingerprint == following.formula_fingerprint
                and previous.formula_fingerprint != cell.formula_fingerprint
            ):
                findings.append(
                    Finding(
                        rule_id="FF006",
                        severity="high",
                        location=location,
                        message=(
                            f"Formula breaks a stable {orientation} pattern between "
                            f"{previous_location[1]} and {next_location[1]}."
                        ),
                        details={
                            "peer_formula_fingerprint": previous.formula_fingerprint,
                            "candidate_formula_fingerprint": cell.formula_fingerprint,
                        },
                    )
                )
                break
    return findings


def _new_coverage_items(
    before: Mapping[CellKey, tuple[str, ...]], after: Mapping[CellKey, tuple[str, ...]]
) -> list[tuple[CellKey, tuple[str, ...]]]:
    """Return per-cell static-analysis coverage items newly present in a candidate."""
    additions: list[tuple[CellKey, tuple[str, ...]]] = []
    for location in sorted(after, key=_location_sort_key):
        previous = {item.casefold() for item in before.get(location, ())}
        new_items = tuple(item for item in after[location] if item.casefold() not in previous)
        if new_items:
            additions.append((location, new_items))
    return additions


def _three_d_reference_scope_changes(
    before: WorkbookSnapshot, after: WorkbookSnapshot
) -> tuple[list[Change], list[Finding]]:
    """Find unchanged 3-D formulas whose workbook tab span now resolves differently."""
    changes: list[Change] = []
    findings: list[Finding] = []
    locations = sorted(
        set(before.three_d_reference_tokens) | set(after.three_d_reference_tokens),
        key=_location_sort_key,
    )
    for location in locations:
        before_cell = before.cells.get(location)
        after_cell = after.cells.get(location)
        if before_cell is None or after_cell is None or before_cell.formula != after_cell.formula:
            continue
        tokens = sorted(
            set(before.three_d_reference_tokens.get(location, ()))
            | set(after.three_d_reference_tokens.get(location, ())),
            key=str.casefold,
        )
        changed_references: list[dict[str, object]] = []
        for token in tokens:
            before_references = resolve_3d_reference(token, before.sheet_order) or ()
            after_references = resolve_3d_reference(token, after.sheet_order) or ()
            before_sheets = [reference.sheet for reference in before_references]
            after_sheets = [reference.sheet for reference in after_references]
            if before_sheets != after_sheets:
                changed_references.append(
                    {
                        "token": token,
                        "before_sheets": before_sheets,
                        "after_sheets": after_sheets,
                    }
                )
        if not changed_references:
            continue
        details = {"references": changed_references}
        changes.append(
            Change(
                "three_d_reference_scope_changed",
                location,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF014",
                "high",
                "Worksheet changes altered the scope of a static 3-D reference.",
                location,
                details=details,
            )
        )
    return changes, findings


def _credential_material_changed(
    before: ProtectionCredentialSnapshot | None,
    after: ProtectionCredentialSnapshot | None,
) -> bool:
    """Compare private verifier signatures without returning verifier material."""
    return (before.signature if before is not None else None) != (
        after.signature if after is not None else None
    )


def _opaque_protection_metadata_changed(
    before: ProtectionOpaqueMetadataSnapshot | None,
    after: ProtectionOpaqueMetadataSnapshot | None,
) -> bool:
    """Compare private opaque metadata fingerprints without exposing content."""
    return (before.signature if before is not None else None) != (
        after.signature if after is not None else None
    )


def _private_external_data_material_changed(
    before: tuple[str | None, ...],
    after: tuple[str | None, ...],
) -> bool:
    """Compare private source or identity fingerprints without returning them."""
    return before != after


def _external_data_source_material_change_categories(
    before: tuple[ExternalDataConnectionSnapshot, ...],
    after: tuple[ExternalDataConnectionSnapshot, ...],
) -> list[str]:
    """Return safe source-material classes that changed on stable connections.

    OOXML connection IDs are the local binding used by QueryTables and pivot
    caches.  Only a one-to-one, non-null ID match is classified; a connection
    add, removal, duplicate ID, or re-numbering remains covered by the broader
    private source-configuration signal without inventing a category.
    """

    def index_by_id(
        connections: tuple[ExternalDataConnectionSnapshot, ...],
    ) -> dict[int, ExternalDataConnectionSnapshot] | None:
        indexed: dict[int, ExternalDataConnectionSnapshot] = {}
        for connection in connections:
            connection_id = connection.connection_id
            if connection_id is None or connection_id in indexed:
                return None
            indexed[connection_id] = connection
        return indexed

    old_connections = index_by_id(before)
    new_connections = index_by_id(after)
    if (
        old_connections is None
        or new_connections is None
        or set(old_connections) != set(new_connections)
    ):
        return []

    categories: set[str] = set()
    for connection_id in sorted(set(old_connections) & set(new_connections)):
        old_signatures = dict(old_connections[connection_id].source_material_signatures)
        new_signatures = dict(new_connections[connection_id].source_material_signatures)
        for category in set(old_signatures) | set(new_signatures):
            if old_signatures.get(category) != new_signatures.get(category):
                categories.add(category)
    return sorted(categories)


def _workbook_control_changes(
    before: WorkbookSnapshot, after: WorkbookSnapshot
) -> tuple[list[Change], list[Finding]]:
    changes: list[Change] = []
    findings: list[Finding] = []

    for title in sorted(set(before.sheets) | set(after.sheets), key=str.casefold):
        old_sheet = before.sheets.get(title)
        new_sheet = after.sheets.get(title)
        if old_sheet is None:
            changes.append(
                Change(
                    "sheet_added",
                    None,
                    "medium",
                    details={"sheet": title, "state": new_sheet.state},
                )
            )
        elif new_sheet is None:
            changes.append(
                Change(
                    "sheet_removed",
                    None,
                    "high",
                    details={"sheet": title, "state": old_sheet.state},
                )
            )
        elif old_sheet.state != new_sheet.state:
            changes.append(
                Change(
                    "sheet_visibility_changed",
                    None,
                    "high",
                    details={"sheet": title, "before": old_sheet.state, "after": new_sheet.state},
                )
            )
            findings.append(
                Finding(
                    "FF007",
                    "high",
                    (
                        "Sheet visibility changed from "
                        f"{old_sheet.state} to {new_sheet.state}: {title}."
                    ),
                    details={"sheet": title, "before": old_sheet.state, "after": new_sheet.state},
                )
                )

    three_d_changes, three_d_findings = _three_d_reference_scope_changes(before, after)
    changes.extend(three_d_changes)
    findings.extend(three_d_findings)

    for name in sorted(set(before.defined_names) | set(after.defined_names), key=str.casefold):
        old_value = before.defined_names.get(name)
        new_value = after.defined_names.get(name)
        if old_value == new_value:
            continue
        changes.append(
            Change(
                "defined_name_changed",
                None,
                "medium",
                details={"name": name, "before": old_value, "after": new_value},
            )
        )
        findings.append(
            Finding(
                "FF008",
                "medium",
                f"Defined name changed: {name}.",
                details={"before": old_value, "after": new_value},
            )
        )

    for name in sorted(set(before.tables) | set(after.tables), key=str.casefold):
        old_table = before.tables.get(name)
        new_table = after.tables.get(name)
        if old_table == new_table:
            continue
        details = {
            "name": name,
            "before": old_table.to_dict() if old_table else None,
            "after": new_table.to_dict() if new_table else None,
        }
        if (
            old_table is not None
            and new_table is not None
            and old_table.calculated_column_formulas
            != new_table.calculated_column_formulas
        ):
            details["calculated_column_formula_material_changed"] = True
        changes.append(
            Change(
                "table_definition_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF013",
                "high",
                f"Excel table definition changed: {name}.",
                details=details,
            )
        )

    before_validations: dict[str, list[DataValidationSnapshot]] = defaultdict(list)
    after_validations: dict[str, list[DataValidationSnapshot]] = defaultdict(list)
    for validation in before.data_validations:
        before_validations[validation.sheet].append(validation)
    for validation in after.data_validations:
        after_validations[validation.sheet].append(validation)
    for sheet in sorted(set(before_validations) | set(after_validations), key=str.casefold):
        old_validations = tuple(before_validations.get(sheet, ()))
        new_validations = tuple(after_validations.get(sheet, ()))
        if old_validations == new_validations:
            continue
        details = {
            "sheet": sheet,
            "before": [validation.to_dict() for validation in old_validations],
            "after": [validation.to_dict() for validation in new_validations],
        }
        changes.append(
            Change(
                "data_validation_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF020",
                "high",
                f"Data-validation controls changed on worksheet: {sheet}.",
                details=details,
            )
        )

    before_conditional_formatting: dict[str, list[ConditionalFormattingSnapshot]] = (
        defaultdict(list)
    )
    after_conditional_formatting: dict[str, list[ConditionalFormattingSnapshot]] = (
        defaultdict(list)
    )
    before_conditional_extensions: dict[
        str, list[ConditionalFormattingExtensionSnapshot]
    ] = defaultdict(list)
    after_conditional_extensions: dict[
        str, list[ConditionalFormattingExtensionSnapshot]
    ] = defaultdict(list)
    for rule in before.conditional_formatting:
        before_conditional_formatting[rule.sheet].append(rule)
    for rule in after.conditional_formatting:
        after_conditional_formatting[rule.sheet].append(rule)
    for extension in before.conditional_formatting_extensions:
        before_conditional_extensions[extension.sheet].append(extension)
    for extension in after.conditional_formatting_extensions:
        after_conditional_extensions[extension.sheet].append(extension)
    conditional_sheets = (
        set(before_conditional_formatting)
        | set(after_conditional_formatting)
        | set(before_conditional_extensions)
        | set(after_conditional_extensions)
    )
    for sheet in sorted(conditional_sheets, key=str.casefold):
        old_rules = tuple(before_conditional_formatting.get(sheet, ()))
        new_rules = tuple(after_conditional_formatting.get(sheet, ()))
        old_extensions = tuple(before_conditional_extensions.get(sheet, ()))
        new_extensions = tuple(after_conditional_extensions.get(sheet, ()))
        if old_rules == new_rules and old_extensions == new_extensions:
            continue
        details = {
            "sheet": sheet,
            "before": {
                "rules": [rule.to_dict() for rule in old_rules],
                "extensions": [extension.to_dict() for extension in old_extensions],
            },
            "after": {
                "rules": [rule.to_dict() for rule in new_rules],
                "extensions": [extension.to_dict() for extension in new_extensions],
            },
        }
        changes.append(
            Change(
                "conditional_formatting_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF021",
                "high",
                f"Conditional-formatting controls changed on worksheet: {sheet}.",
                details=details,
            )
        )

    if before.workbook_protection != after.workbook_protection:
        old_protection = before.workbook_protection
        new_protection = after.workbook_protection
        details: dict[str, object] = {
            "before": old_protection.to_dict() if old_protection is not None else None,
            "after": new_protection.to_dict() if new_protection is not None else None,
        }
        credential_changes = {
            "workbook": _credential_material_changed(
                (
                    old_protection.workbook_credential
                    if old_protection is not None
                    else None
                ),
                (
                    new_protection.workbook_credential
                    if new_protection is not None
                    else None
                ),
            ),
            "revisions": _credential_material_changed(
                (
                    old_protection.revisions_credential
                    if old_protection is not None
                    else None
                ),
                (
                    new_protection.revisions_credential
                    if new_protection is not None
                    else None
                ),
            ),
        }
        if any(credential_changes.values()):
            details["credential_material_changed"] = credential_changes
        if _opaque_protection_metadata_changed(
            old_protection.opaque_metadata if old_protection is not None else None,
            new_protection.opaque_metadata if new_protection is not None else None,
        ):
            details["opaque_metadata_changed"] = True
        changes.append(
            Change(
                "workbook_protection_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF022",
                "high",
                "Workbook protection controls changed.",
                details=details,
            )
        )

    before_sheet_protections = {
        protection.sheet: protection for protection in before.sheet_protections
    }
    after_sheet_protections = {
        protection.sheet: protection for protection in after.sheet_protections
    }
    for sheet in sorted(
        set(before_sheet_protections) | set(after_sheet_protections), key=str.casefold
    ):
        old_protection = before_sheet_protections.get(sheet)
        new_protection = after_sheet_protections.get(sheet)
        if old_protection == new_protection:
            continue
        details = {
            "sheet": sheet,
            "before": old_protection.to_dict() if old_protection is not None else None,
            "after": new_protection.to_dict() if new_protection is not None else None,
        }
        if _credential_material_changed(
            old_protection.credential if old_protection is not None else None,
            new_protection.credential if new_protection is not None else None,
        ):
            details["credential_material_changed"] = True
        if _opaque_protection_metadata_changed(
            old_protection.opaque_metadata if old_protection is not None else None,
            new_protection.opaque_metadata if new_protection is not None else None,
        ):
            details["opaque_metadata_changed"] = True
        changes.append(
            Change(
                "sheet_protection_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF022",
                "high",
                f"Sheet protection controls changed: {sheet}.",
                details=details,
            )
        )

    before_ranges: dict[str, list[object]] = defaultdict(list)
    after_ranges: dict[str, list[object]] = defaultdict(list)
    for protected_range in before.protected_ranges:
        before_ranges[protected_range.sheet].append(protected_range)
    for protected_range in after.protected_ranges:
        after_ranges[protected_range.sheet].append(protected_range)
    for sheet in sorted(set(before_ranges) | set(after_ranges), key=str.casefold):
        old_ranges = tuple(before_ranges.get(sheet, ()))
        new_ranges = tuple(after_ranges.get(sheet, ()))
        if old_ranges == new_ranges:
            continue
        details = {
            "sheet": sheet,
            "before": [protected_range.to_dict() for protected_range in old_ranges],
            "after": [protected_range.to_dict() for protected_range in new_ranges],
        }
        if tuple(item.name_signature for item in old_ranges) != tuple(
            item.name_signature for item in new_ranges
        ):
            details["range_name_material_changed"] = True
        if tuple(item.security_descriptor_signature for item in old_ranges) != tuple(
            item.security_descriptor_signature for item in new_ranges
        ):
            details["security_descriptor_material_changed"] = True
        if tuple(item.credential.signature for item in old_ranges) != tuple(
            item.credential.signature for item in new_ranges
        ):
            details["credential_material_changed"] = True
        if tuple(item.opaque_metadata.signature for item in old_ranges) != tuple(
            item.opaque_metadata.signature for item in new_ranges
        ):
            details["opaque_metadata_changed"] = True
        changes.append(
            Change(
                "protected_range_permissions_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF022",
                "high",
                f"Protected-range permissions changed on worksheet: {sheet}.",
                details=details,
            )
        )

    if before.cell_protection_default != after.cell_protection_default:
        details = {
            "before": (
                before.cell_protection_default.to_dict()
                if before.cell_protection_default is not None
                else None
            ),
            "after": (
                after.cell_protection_default.to_dict()
                if after.cell_protection_default is not None
                else None
            ),
        }
        changes.append(
            Change(
                "cell_protection_default_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF022",
                "high",
                "Default locked/hidden cell protection changed.",
                details=details,
            )
        )

    before_assignments: dict[str, list[object]] = defaultdict(list)
    after_assignments: dict[str, list[object]] = defaultdict(list)
    for assignment in before.cell_protection_assignments:
        before_assignments[assignment.sheet].append(assignment)
    for assignment in after.cell_protection_assignments:
        after_assignments[assignment.sheet].append(assignment)
    for sheet in sorted(
        set(before_assignments) | set(after_assignments), key=str.casefold
    ):
        old_assignments = tuple(before_assignments.get(sheet, ()))
        new_assignments = tuple(after_assignments.get(sheet, ()))
        if old_assignments == new_assignments:
            continue
        details = {
            "sheet": sheet,
            "before": [assignment.to_dict() for assignment in old_assignments],
            "after": [assignment.to_dict() for assignment in new_assignments],
        }
        changes.append(
            Change(
                "cell_protection_assignments_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF022",
                "high",
                f"Direct cell protection assignments changed on worksheet: {sheet}.",
                details=details,
            )
        )

    if before.external_data_refresh_settings != after.external_data_refresh_settings:
        details = {
            "before": before.external_data_refresh_settings.to_dict(),
            "after": after.external_data_refresh_settings.to_dict(),
        }
        changes.append(
            Change(
                "external_data_refresh_settings_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF023",
                "high",
                "Workbook-wide external-data refresh settings changed.",
                details=details,
            )
        )

    if before.external_data_connections != after.external_data_connections:
        old_connections: tuple[ExternalDataConnectionSnapshot, ...] = (
            before.external_data_connections
        )
        new_connections: tuple[ExternalDataConnectionSnapshot, ...] = (
            after.external_data_connections
        )
        details: dict[str, object] = {
            "before": [connection.to_dict() for connection in old_connections],
            "after": [connection.to_dict() for connection in new_connections],
        }
        if _private_external_data_material_changed(
            tuple(connection.identity_signature for connection in old_connections),
            tuple(connection.identity_signature for connection in new_connections),
        ):
            details["identity_material_changed"] = True
        if _private_external_data_material_changed(
            tuple(
                connection.source_configuration_signature
                for connection in old_connections
            ),
            tuple(
                connection.source_configuration_signature
                for connection in new_connections
            ),
        ):
            details["source_configuration_material_changed"] = True
            if categories := _external_data_source_material_change_categories(
                old_connections,
                new_connections,
            ):
                details["source_material_change_categories"] = categories
        if _private_external_data_material_changed(
            tuple(connection.opaque_metadata.signature for connection in old_connections),
            tuple(connection.opaque_metadata.signature for connection in new_connections),
        ):
            details["opaque_metadata_changed"] = True
        changes.append(
            Change(
                "external_data_connections_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF023",
                "high",
                "External-data connection controls changed.",
                details=details,
            )
        )

    if before.query_table_refresh_controls != after.query_table_refresh_controls:
        old_query_tables: tuple[QueryTableRefreshSnapshot, ...] = (
            before.query_table_refresh_controls
        )
        new_query_tables: tuple[QueryTableRefreshSnapshot, ...] = (
            after.query_table_refresh_controls
        )
        details = {
            "before": [control.to_dict() for control in old_query_tables],
            "after": [control.to_dict() for control in new_query_tables],
        }
        if _private_external_data_material_changed(
            tuple(control.identity_signature for control in old_query_tables),
            tuple(control.identity_signature for control in new_query_tables),
        ):
            details["identity_material_changed"] = True
        if _private_external_data_material_changed(
            tuple(control.opaque_metadata.signature for control in old_query_tables),
            tuple(control.opaque_metadata.signature for control in new_query_tables),
        ):
            details["opaque_metadata_changed"] = True
        changes.append(
            Change(
                "query_table_refresh_controls_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF023",
                "high",
                "Query-table refresh controls changed.",
                details=details,
            )
        )

    if before.pivot_cache_refresh_controls != after.pivot_cache_refresh_controls:
        old_pivot_caches: tuple[PivotCacheRefreshSnapshot, ...] = (
            before.pivot_cache_refresh_controls
        )
        new_pivot_caches: tuple[PivotCacheRefreshSnapshot, ...] = (
            after.pivot_cache_refresh_controls
        )
        details = {
            "before": [control.to_dict() for control in old_pivot_caches],
            "after": [control.to_dict() for control in new_pivot_caches],
        }
        if _private_external_data_material_changed(
            tuple(
                control.source_configuration_signature
                for control in old_pivot_caches
            ),
            tuple(
                control.source_configuration_signature
                for control in new_pivot_caches
            ),
        ):
            details["source_configuration_material_changed"] = True
        if _private_external_data_material_changed(
            tuple(control.opaque_metadata.signature for control in old_pivot_caches),
            tuple(control.opaque_metadata.signature for control in new_pivot_caches),
        ):
            details["opaque_metadata_changed"] = True
        changes.append(
            Change(
                "pivot_cache_refresh_controls_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF023",
                "high",
                "Pivot-cache refresh controls changed.",
                details=details,
            )
        )

    if before.external_link_packages != after.external_link_packages:
        old_external_links: ExternalLinkPackageSnapshot = before.external_link_packages
        new_external_links: ExternalLinkPackageSnapshot = after.external_link_packages
        details: dict[str, object] = {
            "before": old_external_links.to_dict(),
            "after": new_external_links.to_dict(),
        }
        if old_external_links.source_signature != new_external_links.source_signature:
            details["source_material_changed"] = True
        if old_external_links.definition_signature != new_external_links.definition_signature:
            details["definition_material_changed"] = True
        if (
            old_external_links.cached_material_signature
            != new_external_links.cached_material_signature
        ):
            details["cached_material_changed"] = True
        if (
            old_external_links.opaque_metadata.signature
            != new_external_links.opaque_metadata.signature
        ):
            details["opaque_metadata_changed"] = True
        changes.append(
            Change(
                "external_link_packages_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF025",
                "high",
                "External-workbook, DDE, or OLE link package controls changed.",
                details=details,
            )
        )

    if before.external_relationships != after.external_relationships:
        old_relationships: ExternalRelationshipSnapshot = before.external_relationships
        new_relationships: ExternalRelationshipSnapshot = after.external_relationships
        details: dict[str, object] = {
            "before": old_relationships.to_dict(),
            "after": new_relationships.to_dict(),
        }
        if (
            old_relationships.relationship_signature
            != new_relationships.relationship_signature
        ):
            details["external_relationship_material_changed"] = True
        if (
            old_relationships.unrecognized_relationship_count
            != new_relationships.unrecognized_relationship_count
            or (
                old_relationships.unrecognized_relationship_count
                or new_relationships.unrecognized_relationship_count
            )
            and (
                old_relationships.relationship_signature
                != new_relationships.relationship_signature
            )
        ):
            details["unrecognized_external_relationship_metadata_changed"] = True
        changes.append(
            Change(
                "external_relationships_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF063",
                "high",
                (
                    "Package-wide external relationships changed; a workbook part may "
                    "now reach a remote resource outside known feature boundaries."
                ),
                details=details,
            )
        )

    if before.power_query != after.power_query:
        old_power_query: PowerQuerySnapshot = before.power_query
        new_power_query: PowerQuerySnapshot = after.power_query
        details: dict[str, object] = {
            "before": old_power_query.to_dict(),
            "after": new_power_query.to_dict(),
        }
        if old_power_query.formula_signature != new_power_query.formula_signature:
            details["formula_material_changed"] = True
        if (
            old_power_query.package_configuration_signature
            != new_power_query.package_configuration_signature
        ):
            details["package_configuration_material_changed"] = True
        if (
            old_power_query.metadata_identity_signature
            != new_power_query.metadata_identity_signature
        ):
            details["metadata_identity_material_changed"] = True
        if (
            old_power_query.metadata_control_signature
            != new_power_query.metadata_control_signature
        ):
            details["metadata_control_material_changed"] = True
        if old_power_query.permission_controls != new_power_query.permission_controls:
            details["permission_controls_changed"] = True
        if (
            old_power_query.permission_binding_count
            != new_power_query.permission_binding_count
        ):
            details["permission_binding_presence_changed"] = True
        if (
            old_power_query.opaque_metadata.signature
            != new_power_query.opaque_metadata.signature
        ):
            details["opaque_metadata_changed"] = True
        changes.append(
            Change(
                "power_query_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF024",
                "high",
                "Power Query formulas or semantic query controls changed.",
                details=details,
            )
        )

    if before.calculation_settings != after.calculation_settings:
        changes.append(
            Change(
                "calculation_settings_changed",
                None,
                "high",
                details={
                    "before": before.calculation_settings,
                    "after": after.calculation_settings,
                },
            )
        )
        findings.append(
            Finding(
                "FF009",
                "high",
                "Workbook calculation settings changed.",
                details={
                    "before": before.calculation_settings,
                    "after": after.calculation_settings,
                },
            )
        )

    if before.workbook_date_system != after.workbook_date_system:
        details: dict[str, object] = {
            "before": before.workbook_date_system.to_dict(),
            "after": after.workbook_date_system.to_dict(),
        }
        if (
            before.workbook_date_system.control_signature
            != after.workbook_date_system.control_signature
        ):
            details["date_system_control_material_changed"] = True
        if (
            before.workbook_date_system.unrecognized_control_count
            != after.workbook_date_system.unrecognized_control_count
        ):
            details["unrecognized_date_system_control_changed"] = True
        changes.append(
            Change(
                "workbook_date_system_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF117",
                "high",
                "Workbook serial-date system controls changed.",
                details=details,
            )
        )

    if before.macro_hash != after.macro_hash:
        changes.append(
            Change(
                "macro_payload_changed",
                None,
                "critical",
                details={"before": before.macro_hash, "after": after.macro_hash},
            )
        )
        findings.append(
            Finding(
                "FF005",
                "critical",
                "VBA macro payload was added, removed, or changed.",
            )
        )
    if before.xlm_macro_sheets != after.xlm_macro_sheets:
        old_xlm_macro_sheets: XlmMacroSheetSnapshot = before.xlm_macro_sheets
        new_xlm_macro_sheets: XlmMacroSheetSnapshot = after.xlm_macro_sheets
        details: dict[str, object] = {
            "before": old_xlm_macro_sheets.to_dict(),
            "after": new_xlm_macro_sheets.to_dict(),
        }
        if (
            old_xlm_macro_sheets.declaration_signature
            != new_xlm_macro_sheets.declaration_signature
        ):
            details["workbook_binding_changed"] = True
        if (
            old_xlm_macro_sheets.program_signature
            != new_xlm_macro_sheets.program_signature
        ):
            details["macro_program_material_changed"] = True
        if (
            old_xlm_macro_sheets.relationship_signature
            != new_xlm_macro_sheets.relationship_signature
        ):
            details["related_part_relationships_changed"] = True
        if (
            old_xlm_macro_sheets.related_part_payload_signature
            != new_xlm_macro_sheets.related_part_payload_signature
        ):
            details["related_part_payload_material_changed"] = True
        changes.append(
            Change(
                "xlm_macro_sheets_changed",
                None,
                "critical",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF026",
                "critical",
                "Excel 4.0 / XLM macro-sheet controls changed.",
                details=details,
            )
        )
    if before.xlm_automatic_macro_bindings != after.xlm_automatic_macro_bindings:
        old_xlm_automatic_macro_bindings: XlmAutomaticMacroBindingSnapshot = (
            before.xlm_automatic_macro_bindings
        )
        new_xlm_automatic_macro_bindings: XlmAutomaticMacroBindingSnapshot = (
            after.xlm_automatic_macro_bindings
        )
        details = {
            "before": old_xlm_automatic_macro_bindings.to_dict(),
            "after": new_xlm_automatic_macro_bindings.to_dict(),
        }
        if (
            old_xlm_automatic_macro_bindings.binding_signature
            != new_xlm_automatic_macro_bindings.binding_signature
        ):
            details["automatic_macro_binding_material_changed"] = True
        changes.append(
            Change(
                "xlm_automatic_macro_bindings_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF076",
                "high",
                "XLM automatic-macro bindings changed.",
                details=details,
            )
        )
    if before.ribbon_customization != after.ribbon_customization:
        old_ribbon: RibbonCustomizationSnapshot = before.ribbon_customization
        new_ribbon: RibbonCustomizationSnapshot = after.ribbon_customization
        details = {
            "before": old_ribbon.to_dict(),
            "after": new_ribbon.to_dict(),
        }
        if old_ribbon.declaration_signature != new_ribbon.declaration_signature:
            details["package_binding_changed"] = True
        if old_ribbon.definition_signature != new_ribbon.definition_signature:
            details["ribbon_definition_material_changed"] = True
        if old_ribbon.relationship_signature != new_ribbon.relationship_signature:
            details["image_relationships_changed"] = True
        changes.append(
            Change(
                "ribbon_customization_changed",
                None,
                "critical",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF027",
                "critical",
                "Office RibbonX customization controls changed.",
                details=details,
            )
        )
    if before.office_web_addins != after.office_web_addins:
        old_addins: OfficeWebAddinSnapshot = before.office_web_addins
        new_addins: OfficeWebAddinSnapshot = after.office_web_addins
        details = {
            "before": old_addins.to_dict(),
            "after": new_addins.to_dict(),
        }
        if old_addins.declaration_signature != new_addins.declaration_signature:
            details["workbook_binding_changed"] = True
        if old_addins.taskpane_signature != new_addins.taskpane_signature:
            details["taskpane_configuration_material_changed"] = True
        if old_addins.web_extension_signature != new_addins.web_extension_signature:
            details["web_extension_definition_material_changed"] = True
        if old_addins.relationship_signature != new_addins.relationship_signature:
            details["related_part_relationships_changed"] = True
        if old_addins.worksheet_binding_signature != new_addins.worksheet_binding_signature:
            details["worksheet_binding_material_changed"] = True
        if old_addins.in_content_signature != new_addins.in_content_signature:
            details["in_content_drawing_binding_changed"] = True
        changes.append(
            Change(
                "office_web_addins_changed",
                None,
                "critical",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF028",
                "critical",
                "Office Web Add-in controls changed.",
                details=details,
            )
        )
    if before.worksheet_embedded_controls != after.worksheet_embedded_controls:
        old_controls: WorksheetEmbeddedControlSnapshot = before.worksheet_embedded_controls
        new_controls: WorksheetEmbeddedControlSnapshot = after.worksheet_embedded_controls
        details = {
            "before": old_controls.to_dict(),
            "after": new_controls.to_dict(),
        }
        if old_controls.declaration_signature != new_controls.declaration_signature:
            details["worksheet_binding_changed"] = True
        if (
            old_controls.control_definition_signature
            != new_controls.control_definition_signature
        ):
            details["worksheet_control_definition_material_changed"] = True
        if (
            old_controls.active_x_definition_signature
            != new_controls.active_x_definition_signature
        ):
            details["active_x_definition_material_changed"] = True
        if (
            old_controls.form_control_property_signature
            != new_controls.form_control_property_signature
        ):
            details["form_control_property_material_changed"] = True
        if (
            old_controls.legacy_vml_definition_signature
            != new_controls.legacy_vml_definition_signature
        ):
            details["legacy_vml_control_definition_material_changed"] = True
        if (
            old_controls.legacy_vml_relationship_signature
            != new_controls.legacy_vml_relationship_signature
        ):
            details["legacy_vml_related_part_relationships_changed"] = True
        if old_controls.relationship_signature != new_controls.relationship_signature:
            details["related_part_relationships_changed"] = True
        if (
            old_controls.related_part_payload_signature
            != new_controls.related_part_payload_signature
        ):
            details["embedded_payload_material_changed"] = True
        changes.append(
            Change(
                "worksheet_embedded_controls_changed",
                None,
                "critical",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF029",
                "critical",
                "Worksheet embedded, legacy VML controls, or OLE objects changed.",
                details=details,
            )
        )
    if before.chart_definitions != after.chart_definitions:
        old_charts: ChartDefinitionSnapshot = before.chart_definitions
        new_charts: ChartDefinitionSnapshot = after.chart_definitions
        details: dict[str, object] = {
            "before": old_charts.to_dict(),
            "after": new_charts.to_dict(),
        }
        if old_charts.declaration_signature != new_charts.declaration_signature:
            details["drawing_binding_changed"] = True
        if old_charts.definition_signature != new_charts.definition_signature:
            details["chart_definition_material_changed"] = True
        if old_charts.cached_data_signature != new_charts.cached_data_signature:
            details["cached_series_material_changed"] = True
        if old_charts.user_shape_signature != new_charts.user_shape_signature:
            details["overlay_shape_material_changed"] = True
        if old_charts.relationship_signature != new_charts.relationship_signature:
            details["related_part_relationships_changed"] = True
        if (
            old_charts.related_part_payload_signature
            != new_charts.related_part_payload_signature
        ):
            details["related_part_payload_material_changed"] = True
        changes.append(
            Change(
                "chart_definitions_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF030",
                "high",
                "Chart definition, cached series data, or overlay shape material changed.",
                details=details,
            )
        )
    if before.pivot_table_definitions != after.pivot_table_definitions:
        old_pivots: PivotTableDefinitionSnapshot = before.pivot_table_definitions
        new_pivots: PivotTableDefinitionSnapshot = after.pivot_table_definitions
        details: dict[str, object] = {
            "before": old_pivots.to_dict(),
            "after": new_pivots.to_dict(),
        }
        if old_pivots.declaration_signature != new_pivots.declaration_signature:
            details["pivot_table_binding_changed"] = True
        if old_pivots.layout_signature != new_pivots.layout_signature:
            details["pivot_table_layout_material_changed"] = True
        if (
            old_pivots.cache_definition_signature
            != new_pivots.cache_definition_signature
        ):
            details["pivot_cache_definition_material_changed"] = True
        if (
            old_pivots.cached_shared_item_signature
            != new_pivots.cached_shared_item_signature
        ):
            details["cached_shared_item_material_changed"] = True
        if old_pivots.relationship_signature != new_pivots.relationship_signature:
            details["related_part_relationships_changed"] = True
        if (
            old_pivots.cache_record_payload_signature
            != new_pivots.cache_record_payload_signature
        ):
            details["cache_record_payload_material_changed"] = True
        changes.append(
            Change(
                "pivot_table_definitions_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF031",
                "high",
                "PivotTable view, cache-schema, shared-item, or cached-record material changed.",
                details=details,
            )
        )
    if before.slicer_timeline_caches != after.slicer_timeline_caches:
        old_filters: SlicerTimelineCacheSnapshot = before.slicer_timeline_caches
        new_filters: SlicerTimelineCacheSnapshot = after.slicer_timeline_caches
        details: dict[str, object] = {
            "before": old_filters.to_dict(),
            "after": new_filters.to_dict(),
        }
        if old_filters.declaration_signature != new_filters.declaration_signature:
            details["workbook_cache_binding_changed"] = True
        if (
            old_filters.slicer_definition_signature
            != new_filters.slicer_definition_signature
        ):
            details["slicer_filter_state_or_definition_material_changed"] = True
        if (
            old_filters.timeline_definition_signature
            != new_filters.timeline_definition_signature
        ):
            details["timeline_filter_state_or_definition_material_changed"] = True
        if old_filters.relationship_signature != new_filters.relationship_signature:
            details["related_part_relationships_changed"] = True
        changes.append(
            Change(
                "slicer_timeline_cache_definitions_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF032",
                "high",
                "Slicer or Timeline cache filter state or definition changed.",
                details=details,
            )
        )
    if before.power_pivot_data_model != after.power_pivot_data_model:
        old_data_model: PowerPivotDataModelSnapshot = before.power_pivot_data_model
        new_data_model: PowerPivotDataModelSnapshot = after.power_pivot_data_model
        details: dict[str, object] = {
            "before": old_data_model.to_dict(),
            "after": new_data_model.to_dict(),
        }
        if old_data_model.declaration_signature != new_data_model.declaration_signature:
            details["workbook_data_model_declaration_changed"] = True
        if old_data_model.relationship_signature != new_data_model.relationship_signature:
            details["related_data_model_relationships_changed"] = True
        if old_data_model.payload_signature != new_data_model.payload_signature:
            details["embedded_data_model_payload_changed"] = True
        changes.append(
            Change(
                "power_pivot_data_model_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF033",
                "high",
                "Embedded Power Pivot/Data Model definition or payload changed.",
                details=details,
            )
        )
    if before.what_if_data_tables != after.what_if_data_tables:
        old_data_tables: WhatIfDataTableSnapshot = before.what_if_data_tables
        new_data_tables: WhatIfDataTableSnapshot = after.what_if_data_tables
        details: dict[str, object] = {
            "before": old_data_tables.to_dict(),
            "after": new_data_tables.to_dict(),
        }
        if old_data_tables.definition_signature != new_data_tables.definition_signature:
            details["data_table_definition_material_changed"] = True
        changes.append(
            Change(
                "what_if_data_tables_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF034",
                "high",
                "What-If Data Table sensitivity definition or control changed.",
                details=details,
            )
        )
    if before.scenario_manager != after.scenario_manager:
        old_scenarios: ScenarioManagerSnapshot = before.scenario_manager
        new_scenarios: ScenarioManagerSnapshot = after.scenario_manager
        details: dict[str, object] = {
            "before": old_scenarios.to_dict(),
            "after": new_scenarios.to_dict(),
        }
        if old_scenarios.definition_signature != new_scenarios.definition_signature:
            details["scenario_definition_material_changed"] = True
        changes.append(
            Change(
                "scenario_manager_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF035",
                "high",
                "Excel Scenario Manager definition or stored input set changed.",
                details=details,
            )
        )
    if before.filter_visibility_controls != after.filter_visibility_controls:
        old_controls: FilterVisibilitySnapshot = before.filter_visibility_controls
        new_controls: FilterVisibilitySnapshot = after.filter_visibility_controls
        details: dict[str, object] = {
            "before": old_controls.to_dict(),
            "after": new_controls.to_dict(),
        }
        if old_controls.definition_signature != new_controls.definition_signature:
            details["filter_visibility_definition_material_changed"] = True
        changes.append(
            Change(
                "filter_visibility_controls_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF036",
                "high",
                "Worksheet/table filter, sort, or row/column-visibility control changed.",
                details=details,
            )
        )
    if before.ignored_error_controls != after.ignored_error_controls:
        old_controls: IgnoredErrorSnapshot = before.ignored_error_controls
        new_controls: IgnoredErrorSnapshot = after.ignored_error_controls
        details: dict[str, object] = {
            "before": old_controls.to_dict(),
            "after": new_controls.to_dict(),
        }
        if old_controls.definition_signature != new_controls.definition_signature:
            details["ignored_error_definition_material_changed"] = True
        changes.append(
            Change(
                "ignored_error_controls_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF037",
                "high",
                "Excel ignored-error controls changed; review warnings may be suppressed "
                "or restored.",
                details=details,
            )
        )
    if before.named_sheet_views != after.named_sheet_views:
        old_views: NamedSheetViewSnapshot = before.named_sheet_views
        new_views: NamedSheetViewSnapshot = after.named_sheet_views
        details: dict[str, object] = {
            "before": old_views.to_dict(),
            "after": new_views.to_dict(),
        }
        if old_views.definition_signature != new_views.definition_signature:
            details["named_sheet_view_definition_material_changed"] = True
        changes.append(
            Change(
                "named_sheet_views_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF038",
                "high",
                "Excel Named Sheet View controls changed; alternate filter or sort views "
                "may show a different report.",
                details=details,
            )
        )
    if before.custom_workbook_views != after.custom_workbook_views:
        old_views: CustomWorkbookViewSnapshot = before.custom_workbook_views
        new_views: CustomWorkbookViewSnapshot = after.custom_workbook_views
        details: dict[str, object] = {
            "before": old_views.to_dict(),
            "after": new_views.to_dict(),
        }
        if old_views.definition_signature != new_views.definition_signature:
            details["custom_workbook_view_definition_material_changed"] = True
        if old_views.unrecognized_signature != new_views.unrecognized_signature:
            details["unrecognized_custom_workbook_view_metadata_changed"] = True
        changes.append(
            Change(
                "custom_workbook_views_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF060",
                "high",
                "Legacy Excel Custom Views changed; saved alternate display, visibility, "
                "filter, or print settings may show a different workbook.",
                details=details,
            )
        )
    if before.table_style_controls != after.table_style_controls:
        old_controls: TableStyleControlsSnapshot = before.table_style_controls
        new_controls: TableStyleControlsSnapshot = after.table_style_controls
        details: dict[str, object] = {
            "before": old_controls.to_dict(),
            "after": new_controls.to_dict(),
        }
        if old_controls.definition_signature != new_controls.definition_signature:
            details["table_style_definition_material_changed"] = True
        if old_controls.unrecognized_signature != new_controls.unrecognized_signature:
            details["unrecognized_table_style_metadata_changed"] = True
        changes.append(
            Change(
                "table_style_controls_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF061",
                "high",
                "Excel Table Style controls changed; headers, totals, data areas, borders, "
                "banding, or emphasized columns may present a different review surface.",
                details=details,
            )
        )
    if before.shared_workbook_revisions != after.shared_workbook_revisions:
        old_revisions: SharedWorkbookRevisionSnapshot = before.shared_workbook_revisions
        new_revisions: SharedWorkbookRevisionSnapshot = after.shared_workbook_revisions
        details: dict[str, object] = {
            "before": old_revisions.to_dict(),
            "after": new_revisions.to_dict(),
        }
        if old_revisions.header_signature != new_revisions.header_signature:
            details["revision_header_material_changed"] = True
        if old_revisions.log_signature != new_revisions.log_signature:
            details["revision_log_material_changed"] = True
        if old_revisions.relationship_signature != new_revisions.relationship_signature:
            details["revision_relationship_material_changed"] = True
        if old_revisions.unrecognized_signature != new_revisions.unrecognized_signature:
            details["unrecognized_shared_workbook_revision_metadata_changed"] = True
        changes.append(
            Change(
                "shared_workbook_revisions_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF062",
                "high",
                "Legacy shared-workbook revision history changed; prior values, audit "
                "trail, tracking, or conflict-resolution controls may differ.",
                details=details,
            )
        )
    if before.number_format_controls != after.number_format_controls:
        old_controls: NumberFormatSnapshot = before.number_format_controls
        new_controls: NumberFormatSnapshot = after.number_format_controls
        details: dict[str, object] = {
            "before": old_controls.to_dict(),
            "after": new_controls.to_dict(),
        }
        if old_controls.definition_signature != new_controls.definition_signature:
            details["number_format_definition_material_changed"] = True
        changes.append(
            Change(
                "number_format_controls_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF039",
                "high",
                "Cell number-format controls changed; displayed values may be hidden, "
                "scaled, or reinterpreted.",
                details=details,
            )
        )
    if before.font_controls != after.font_controls:
        old_controls: FontSnapshot = before.font_controls
        new_controls: FontSnapshot = after.font_controls
        details = {
            "before": old_controls.to_dict(),
            "after": new_controls.to_dict(),
        }
        if old_controls.definition_signature != new_controls.definition_signature:
            details["font_definition_material_changed"] = True
        changes.append(
            Change(
                "font_controls_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF040",
                "high",
                "Cell font controls changed; values or warnings may be made less visible "
                "or misleading.",
                details=details,
            )
        )
    if before.fill_controls != after.fill_controls:
        old_controls: FillSnapshot = before.fill_controls
        new_controls: FillSnapshot = after.fill_controls
        details = {
            "before": old_controls.to_dict(),
            "after": new_controls.to_dict(),
        }
        if old_controls.definition_signature != new_controls.definition_signature:
            details["fill_definition_material_changed"] = True
        changes.append(
            Change(
                "fill_controls_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF041",
                "high",
                "Cell fill controls changed; values, warnings, or visual classifications "
                "may be made less visible or misleading.",
                details=details,
            )
        )
    if before.workbook_theme != after.workbook_theme:
        old_theme: WorkbookThemeSnapshot = before.workbook_theme
        new_theme: WorkbookThemeSnapshot = after.workbook_theme
        details = {
            "before": old_theme.to_dict(),
            "after": new_theme.to_dict(),
        }
        if old_theme.definition_signature != new_theme.definition_signature:
            details["theme_definition_material_changed"] = True
        if old_theme.image_payload_signature != new_theme.image_payload_signature:
            details["theme_image_payload_changed"] = True
        if old_theme.relationship_signature != new_theme.relationship_signature:
            details["theme_relationships_changed"] = True
        if (
            old_theme.unrecognized_theme_count != new_theme.unrecognized_theme_count
            or (
                (
                    old_theme.unrecognized_theme_count
                    or new_theme.unrecognized_theme_count
                )
                and (
                    old_theme.definition_signature != new_theme.definition_signature
                    or (
                        old_theme.image_payload_signature
                        != new_theme.image_payload_signature
                    )
                    or (
                        old_theme.relationship_signature
                        != new_theme.relationship_signature
                    )
                )
            )
        ):
            details["unrecognized_theme_metadata_changed"] = True
        changes.append(
            Change(
                "workbook_theme_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF053",
                "high",
                (
                    "Workbook theme controls changed; themed cell, chart, or drawing "
                    "appearance may have been altered outside ordinary cells."
                ),
                details=details,
            )
        )
    if before.alignment_controls != after.alignment_controls:
        old_controls: AlignmentSnapshot = before.alignment_controls
        new_controls: AlignmentSnapshot = after.alignment_controls
        details = {
            "before": old_controls.to_dict(),
            "after": new_controls.to_dict(),
        }
        if old_controls.definition_signature != new_controls.definition_signature:
            details["alignment_definition_material_changed"] = True
        if (
            old_controls.unrecognized_alignment_count
            != new_controls.unrecognized_alignment_count
            or (
                (
                    old_controls.unrecognized_alignment_count
                    or new_controls.unrecognized_alignment_count
                )
                and (
                    old_controls.definition_signature
                    != new_controls.definition_signature
                )
            )
        ):
            details["unrecognized_alignment_metadata_changed"] = True
        changes.append(
            Change(
                "cell_alignment_controls_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF054",
                "high",
                (
                    "Cell alignment controls changed; values, warnings, or visual "
                    "classifications may be repositioned, rotated, wrapped, shrunk, "
                    "or made less legible."
                ),
                details=details,
            )
        )
    if before.border_controls != after.border_controls:
        old_controls: BorderSnapshot = before.border_controls
        new_controls: BorderSnapshot = after.border_controls
        details = {
            "before": old_controls.to_dict(),
            "after": new_controls.to_dict(),
        }
        if old_controls.definition_signature != new_controls.definition_signature:
            details["border_definition_material_changed"] = True
        if (
            old_controls.unrecognized_border_count
            != new_controls.unrecognized_border_count
            or (
                (
                    old_controls.unrecognized_border_count
                    or new_controls.unrecognized_border_count
                )
                and (
                    old_controls.unrecognized_signature
                    != new_controls.unrecognized_signature
                )
            )
        ):
            details["unrecognized_border_metadata_changed"] = True
        changes.append(
            Change(
                "cell_border_controls_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF057",
                "high",
                (
                    "Cell border controls changed; report boundaries, totals, warnings, "
                    "or printed presentation may be altered."
                ),
                details=details,
            )
        )
    if before.worksheet_dimension_controls != after.worksheet_dimension_controls:
        old_controls: WorksheetDimensionSnapshot = before.worksheet_dimension_controls
        new_controls: WorksheetDimensionSnapshot = after.worksheet_dimension_controls
        details = {
            "before": old_controls.to_dict(),
            "after": new_controls.to_dict(),
        }
        if old_controls.definition_signature != new_controls.definition_signature:
            details["worksheet_dimension_definition_material_changed"] = True
        if (
            old_controls.unrecognized_dimension_count
            != new_controls.unrecognized_dimension_count
            or (
                (
                    old_controls.unrecognized_dimension_count
                    or new_controls.unrecognized_dimension_count
                )
                and (
                    old_controls.unrecognized_signature
                    != new_controls.unrecognized_signature
                )
            )
        ):
            details["unrecognized_worksheet_dimension_metadata_changed"] = True
        changes.append(
            Change(
                "worksheet_dimension_controls_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF058",
                "high",
                (
                    "Worksheet dimension controls changed; wrapped content may be "
                    "truncated, report framing altered, or automatic pagination shifted."
                ),
                details=details,
            )
        )
    if before.worksheet_display_controls != after.worksheet_display_controls:
        old_controls: WorksheetDisplaySnapshot = before.worksheet_display_controls
        new_controls: WorksheetDisplaySnapshot = after.worksheet_display_controls
        details = {
            "before": old_controls.to_dict(),
            "after": new_controls.to_dict(),
        }
        if old_controls.definition_signature != new_controls.definition_signature:
            details["worksheet_display_definition_material_changed"] = True
        if (
            old_controls.unrecognized_display_control_count
            != new_controls.unrecognized_display_control_count
            or (
                (
                    old_controls.unrecognized_display_control_count
                    or new_controls.unrecognized_display_control_count
                )
                and (
                    old_controls.definition_signature
                    != new_controls.definition_signature
                )
            )
        ):
            details["unrecognized_worksheet_display_metadata_changed"] = True
        changes.append(
            Change(
                "worksheet_display_controls_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF055",
                "high",
                (
                    "Worksheet display controls changed; zeroes, formulas, gridline "
                    "colour, headers, page whitespace, rulers, outline symbols, or "
                    "saved views and panes may alter the reviewer-visible surface."
                ),
                details=details,
            )
        )
    if (
        before.worksheet_print_layout_controls
        != after.worksheet_print_layout_controls
    ):
        old_controls: WorksheetPrintLayoutSnapshot = (
            before.worksheet_print_layout_controls
        )
        new_controls: WorksheetPrintLayoutSnapshot = (
            after.worksheet_print_layout_controls
        )
        details = {
            "before": old_controls.to_dict(),
            "after": new_controls.to_dict(),
        }
        if old_controls.definition_signature != new_controls.definition_signature:
            details["worksheet_print_layout_definition_material_changed"] = True
        if (
            old_controls.unrecognized_print_layout_count
            != new_controls.unrecognized_print_layout_count
            or old_controls.unrecognized_signature
            != new_controls.unrecognized_signature
        ):
            details["unrecognized_worksheet_print_layout_metadata_changed"] = True
        changes.append(
            Change(
                "worksheet_print_layout_controls_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF056",
                "high",
                (
                    "Worksheet print-layout controls changed; print areas, repeated "
                    "titles, margins, page setup, headers or footers, or manual page "
                    "breaks may alter the saved print surface."
                ),
                details=details,
            )
        )
    new_parser_warnings = sorted(set(after.parser_warnings) - set(before.parser_warnings))
    if new_parser_warnings:
        changes.append(
            Change(
                "parser_coverage_warning_added",
                None,
                "medium",
                details={"warnings": new_parser_warnings},
            )
        )
        findings.append(
            Finding(
                "FF010",
                "medium",
                (
                    "Candidate contains unsupported workbook features; "
                    "inspection coverage may be incomplete."
                ),
                details={"warnings": new_parser_warnings},
            )
        )
    return changes, findings


def _array_formula_semantics_changes(
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
    *,
    state_budget: ChangeAnalysisBudget,
) -> tuple[list[Change], list[Finding]]:
    """Report changes Excel would otherwise hide behind identical formula text.

    The same formula text can run as an ordinary scalar formula, a legacy CSE
    array formula with a fixed output range, or a dynamic array. Dynamic result
    extents are intentionally excluded because they can change at recalc time.
    """
    changes: list[Change] = []
    findings: list[Finding] = []
    known_modes = {"absent", "value", "ordinary", "legacy_cse", "dynamic"}

    if (
        before.array_formula_metadata_complete
        != after.array_formula_metadata_complete
        or before.array_formula_metadata_coverage_signature
        != after.array_formula_metadata_coverage_signature
    ):
        details: dict[str, object] = {
            "before_array_formula_metadata_complete": (
                before.array_formula_metadata_complete
            ),
            "after_array_formula_metadata_complete": (
                after.array_formula_metadata_complete
            ),
        }
        if (
            before.array_formula_metadata_coverage_signature
            != after.array_formula_metadata_coverage_signature
        ):
            details["array_formula_metadata_coverage_material_changed"] = True
        changes.append(
            Change(
                "array_formula_metadata_coverage_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF018",
                "high",
                (
                    "Raw OOXML array-formula metadata coverage changed; dynamic-array "
                    "or fixed-CSE classification may be incomplete."
                ),
                details=details,
            )
        )

    def formula_mode(cell: CellSnapshot | None) -> str:
        if cell is None:
            return "absent"
        if not cell.is_formula:
            return "value"
        return cell.array_formula_kind or "ordinary"

    def details_with_impact(location: CellKey, details: dict[str, object]) -> tuple[
        dict[str, object], int, tuple[CellKey, ...]
    ]:
        impact_analysis = analyze_downstream_impact(
            location,
            before,
            after,
            state_budget=state_budget,
        )
        sampled_impacts = tuple(
            sorted(impact_analysis.impacted, key=_location_sort_key)[:_IMPACT_SAMPLE_SIZE]
        )
        if sampled_impacts:
            details["impact_paths"] = _serialise_impact_paths(
                sampled_impacts, impact_analysis.paths
            )
        if impact_analysis.truncated:
            details["impact_truncated_at"] = _IMPACT_NODE_LIMIT
        return details, len(impact_analysis.impacted), sampled_impacts

    locations = sorted(set(before.cells) | set(after.cells), key=_location_sort_key)
    for location in locations:
        old_cell = before.cells.get(location)
        new_cell = after.cells.get(location)
        old_mode = formula_mode(old_cell)
        new_mode = formula_mode(new_cell)
        if old_mode not in known_modes or new_mode not in known_modes:
            continue
        if old_mode != new_mode and {old_mode, new_mode} & {"legacy_cse", "dynamic"}:
            details, impact_count, impacted_cells = details_with_impact(location, {
                "before": {
                    "mode": old_mode,
                    "output_range": old_cell.array_formula_ref if old_cell else None,
                },
                "after": {
                    "mode": new_mode,
                    "output_range": new_cell.array_formula_ref if new_cell else None,
                },
            })
            changes.append(
                Change(
                    "array_formula_mode_changed",
                    location,
                    "high",
                    impact_count=impact_count,
                    impacted_cells=impacted_cells,
                    details=details,
                )
            )
            findings.append(
                Finding(
                    "FF018",
                    "high",
                    (
                        "Formula changed to or from legacy CSE or dynamic array semantics."
                    ),
                    location,
                    details=details,
                )
            )
            continue
        if (
            old_mode == "legacy_cse"
            and new_mode == "legacy_cse"
            and old_cell is not None
            and new_cell is not None
            and old_cell.array_formula_ref != new_cell.array_formula_ref
        ):
            details, impact_count, impacted_cells = details_with_impact(location, {
                "before_output_range": old_cell.array_formula_ref,
                "after_output_range": new_cell.array_formula_ref,
            })
            changes.append(
                Change(
                    "legacy_array_output_range_changed",
                    location,
                    "high",
                    impact_count=impact_count,
                    impacted_cells=impacted_cells,
                    details=details,
                )
            )
            findings.append(
                Finding(
                    "FF018",
                    "high",
                    "Legacy CSE array formula fixed output range changed.",
                    location,
                    details=details,
                )
            )
    return changes, findings


def _dynamic_array_output_reference_changes(
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
    *,
    state_budget: ChangeAnalysisBudget,
) -> tuple[list[Change], list[Finding]]:
    """Report formulas newly intersecting an observed dynamic spill member.

    The relationship is deliberately keyed by consumer and anchor, rather than
    its observed range text. A dynamic array can resize during recalc, so a
    changed cached extent alone must not act like a fixed-range semantic change.
    """
    changes: list[Change] = []
    findings: list[Finding] = []
    for location in sorted(after.dynamic_array_output_references, key=_location_sort_key):
        previous_anchors = {
            reference.anchor
            for reference in before.dynamic_array_output_references.get(location, ())
        }
        new_references = tuple(
            reference
            for reference in after.dynamic_array_output_references[location]
            if reference.anchor not in previous_anchors
        )
        if not new_references:
            continue
        impact_analysis = analyze_downstream_impact(
            location,
            before,
            after,
            state_budget=state_budget,
        )
        sampled_impacts = tuple(
            sorted(impact_analysis.impacted, key=_location_sort_key)[:_IMPACT_SAMPLE_SIZE]
        )
        details: dict[str, object] = {
            "references": [reference.to_dict() for reference in new_references],
        }
        if sampled_impacts:
            details["impact_paths"] = _serialise_impact_paths(
                sampled_impacts, impact_analysis.paths
            )
        if impact_analysis.truncated:
            details["impact_truncated_at"] = _IMPACT_NODE_LIMIT
        changes.append(
            Change(
                "dynamic_array_output_reference_added",
                location,
                "medium",
                impact_count=len(impact_analysis.impacted),
                impacted_cells=sampled_impacts,
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF019",
                "medium",
                (
                    "Formula newly intersects a non-anchor member of an observed "
                    "dynamic-array spill; the observed output extent can resize."
                ),
                location,
                details=details,
            )
        )
    return changes, findings


def _formula_cached_result_entry_map(
    snapshot: FormulaCachedResultSnapshot,
) -> dict[CellKey, tuple[tuple[str, str], ...]]:
    """Index private result digests without returning them in review evidence."""
    entries: dict[CellKey, list[tuple[str, str]]] = defaultdict(list)
    for entry in snapshot.entries:
        entries[entry.location].append((entry.result_type, entry.result_signature))
    return {
        location: tuple(sorted(values))
        for location, values in entries.items()
    }


def _formula_cached_result_changes(
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
    semantic_cell_changes: Iterable[Change],
    *,
    state_budget: ChangeAnalysisBudget,
) -> tuple[list[Change], list[Finding]]:
    """Flag result-cache changes that static visible edits cannot explain.

    Stored formula results are not exposed by the normal formula reader. A
    calculation following an ordinary input or formula edit can legitimately
    update them, so explicit downstream dependencies suppress those cases.
    FormulaFence deliberately does not evaluate formulas or infer dynamic,
    volatile, or external inputs: a remaining change is review evidence, not a
    claim that the result is mathematically incorrect.
    """
    old_results = before.formula_cached_results
    new_results = after.formula_cached_results
    if old_results == new_results:
        return [], []

    old_entries = _formula_cached_result_entry_map(old_results)
    new_entries = _formula_cached_result_entry_map(new_results)
    changed_locations = {
        location
        for location in set(old_entries) | set(new_entries)
        if old_entries.get(location) != new_entries.get(location)
    }
    unexplained_locations = set(changed_locations)
    changes_with_locations = tuple(
        change for change in semantic_cell_changes if change.location is not None
    )

    # A changed formula owns its own cache. A value/formula edit can also
    # explain formula caches reachable through FormulaFence's static graph.
    for change in changes_with_locations:
        old_cell = change.before
        new_cell = change.after
        if (old_cell is not None and old_cell.is_formula) or (
            new_cell is not None and new_cell.is_formula
        ):
            unexplained_locations.discard(change.location)

    for change in changes_with_locations:
        if not unexplained_locations:
            break
        impact = analyze_downstream_impact(
            change.location,
            before,
            after,
            state_budget=state_budget,
        )
        unexplained_locations.difference_update(impact.impacted)

    unrecognized_metadata_changed = (
        old_results.unrecognized_cached_result_count
        != new_results.unrecognized_cached_result_count
    )
    if not unexplained_locations and not unrecognized_metadata_changed:
        return [], []

    details: dict[str, object] = {
        "before": old_results.to_dict(),
        "after": new_results.to_dict(),
        "unexplained_cached_result_change_count": len(unexplained_locations),
    }
    if old_results.definition_signature != new_results.definition_signature:
        details["cached_result_material_changed"] = True
    if unrecognized_metadata_changed:
        details["unrecognized_cached_result_metadata_changed"] = True
    change = Change(
        "formula_cached_result_changed",
        None,
        "high",
        details=details,
    )
    finding = Finding(
        "FF042",
        "high",
        (
            "Stored formula results changed without a formula or statically visible "
            "precedent change."
        ),
        details=details,
    )
    return [change], [finding]


def _formula_external_action_changes(
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
    static_input_change_locations: set[CellKey],
) -> tuple[list[Change], list[Finding]]:
    """Flag changed stored action calls and their statically visible inputs.

    Formula arguments are never evaluated. Instead, a normal cell edit is
    treated as action-relevant only when FormulaFence's existing static
    dependency graph can reach a known action formula in either snapshot.
    Dynamic or unresolved arguments remain ordinary parser-coverage limits.
    """
    old_actions: FormulaExternalActionSnapshot = before.formula_external_actions
    new_actions: FormulaExternalActionSnapshot = after.formula_external_actions
    if old_actions == new_actions and not static_input_change_locations:
        return [], []

    details: dict[str, object] = {
        "before": old_actions.to_dict(),
        "after": new_actions.to_dict(),
    }
    if old_actions.action_signature != new_actions.action_signature:
        details["formula_external_action_material_changed"] = True
    if old_actions.definition_signature != new_actions.definition_signature:
        details["formula_external_action_definition_material_changed"] = True
    if static_input_change_locations:
        details["formula_external_action_static_input_changed"] = True
        details["formula_external_action_static_input_change_count"] = len(
            static_input_change_locations
        )
    change = Change(
        "formula_external_actions_changed",
        None,
        "high",
        details=details,
    )
    finding = Finding(
        "FF064",
        "high",
        (
            "Formula external-action or data-provider functions, a relevant "
            "formula-defined name, or a statically visible input changed; a formula "
            "may now redirect a reviewer, request content, query a cube, or invoke a "
            "real-time data provider."
        ),
        details=details,
    )
    return [change], [finding]


def _formula_dde_link_changes(
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
    static_input_change_locations: set[CellKey],
) -> tuple[list[Change], list[Finding]]:
    """Flag direct DDE formula material and statically visible input changes.

    FormulaFence compares only its private lexical DDE inventory and existing
    static dependency graph.  It never evaluates a formula, starts or contacts
    a DDE server, resolves an application/topic/item, or exposes an endpoint.
    Raw external-link packages remain a distinct package-level boundary.
    """
    old_links: FormulaDdeLinkSnapshot = before.formula_dde_links
    new_links: FormulaDdeLinkSnapshot = after.formula_dde_links
    if old_links == new_links and not static_input_change_locations:
        return [], []

    details: dict[str, object] = {
        "before": old_links.to_dict(),
        "after": new_links.to_dict(),
    }
    if old_links.invocation_signature != new_links.invocation_signature:
        details["formula_dde_link_invocation_material_changed"] = True
    if old_links.definition_signature != new_links.definition_signature:
        details["formula_dde_link_definition_material_changed"] = True
    if static_input_change_locations:
        details["formula_dde_link_static_input_changed"] = True
        details["formula_dde_link_static_input_change_count"] = len(
            static_input_change_locations
        )
    change = Change(
        "formula_dde_links_changed",
        None,
        "high",
        details=details,
    )
    finding = Finding(
        "FF074",
        "high",
        (
            "A direct DDE-style formula link, a relevant formula-defined name, or "
            "a statically visible input changed; depending on local Excel security "
            "settings, the workbook may communicate with or launch a DDE server."
        ),
        details=details,
    )
    return [change], [finding]


def _python_in_excel_changes(
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
    static_input_change_locations: set[CellKey],
) -> tuple[list[Change], list[Finding]]:
    """Flag stored Python code, PY bindings, and statically visible inputs.

    Python-in-Excel code is stored separately from the formula placeholder.
    FormulaFence compares that package material privately and uses its existing
    static graph only to identify ordinary cell edits that can reach a PY call.
    It never parses code as Python, evaluates a formula, or contacts the cloud
    runtime. Dynamic or unresolved inputs remain normal parser-coverage limits.
    """
    old_python: PythonInExcelSnapshot = before.python_in_excel
    new_python: PythonInExcelSnapshot = after.python_in_excel
    if old_python == new_python and not static_input_change_locations:
        return [], []

    details: dict[str, object] = {
        "before": old_python.to_dict(),
        "after": new_python.to_dict(),
    }
    if old_python.definition_signature != new_python.definition_signature:
        details["python_in_excel_definition_changed"] = True
    if old_python.formula_signature != new_python.formula_signature:
        details["python_in_excel_formula_binding_changed"] = True
    if static_input_change_locations:
        details["python_in_excel_static_input_changed"] = True
        details["python_in_excel_static_input_change_count"] = len(
            static_input_change_locations
        )
    change = Change(
        "python_in_excel_changed",
        None,
        "high",
        details=details,
    )
    finding = Finding(
        "FF065",
        "high",
        (
            "Python-in-Excel code, a PY formula binding, or a statically visible "
            "input changed; the workbook may now run different Python code in the "
            "Microsoft Cloud."
        ),
        details=details,
    )
    return [change], [finding]


def _office_custom_function_changes(
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
    static_input_change_locations: set[CellKey],
) -> tuple[list[Change], list[Finding]]:
    """Flag namespaced custom-function candidates and static input changes.

    The workbook alone does not carry an Office Add-in's manifest or JavaScript
    runtime. FormulaFence therefore guards only a stored, namespaced callable
    candidate and normal cell edits that can reach one through its static graph;
    it never resolves a candidate to an add-in, loads an add-in, evaluates a
    formula, or makes a web request.
    """
    old_functions: OfficeCustomFunctionSnapshot = before.office_custom_functions
    new_functions: OfficeCustomFunctionSnapshot = after.office_custom_functions
    if old_functions == new_functions and not static_input_change_locations:
        return [], []

    details: dict[str, object] = {
        "before": old_functions.to_dict(),
        "after": new_functions.to_dict(),
    }
    if old_functions.call_signature != new_functions.call_signature:
        details["office_custom_function_material_changed"] = True
    if old_functions.definition_signature != new_functions.definition_signature:
        details["office_custom_function_definition_changed"] = True
    if static_input_change_locations:
        details["office_custom_function_static_input_changed"] = True
        details["office_custom_function_static_input_change_count"] = len(
            static_input_change_locations
        )
    change = Change(
        "office_custom_functions_changed",
        None,
        "high",
        details=details,
    )
    finding = Finding(
        "FF066",
        "high",
        (
            "A namespaced custom-function call or a statically visible input changed; "
            "a formula may now invoke a different Office Add-in runtime."
        ),
        details=details,
    )
    return [change], [finding]


def _unqualified_runtime_function_changes(
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
    static_input_change_locations: set[CellKey],
) -> tuple[list[Change], list[Finding]]:
    """Flag unknown bare callable candidates and static input changes.

    A workbook does not identify the provider for an unknown unqualified call.
    FormulaFence therefore compares the stored candidate surface and relevant
    formula-defined-name chain privately, plus ordinary edits that can reach a
    candidate through the static dependency graph. It does not resolve VBA,
    COM/Automation add-ins, XLLs, or any registered runtime; execute a formula;
    or inspect the host environment.
    """
    old_functions: UnqualifiedRuntimeFunctionSnapshot = (
        before.unqualified_runtime_functions
    )
    new_functions: UnqualifiedRuntimeFunctionSnapshot = (
        after.unqualified_runtime_functions
    )
    if old_functions == new_functions and not static_input_change_locations:
        return [], []

    details: dict[str, object] = {
        "before": old_functions.to_dict(),
        "after": new_functions.to_dict(),
    }
    if old_functions.call_signature != new_functions.call_signature:
        details["unqualified_runtime_function_material_changed"] = True
    if old_functions.definition_signature != new_functions.definition_signature:
        details["unqualified_runtime_function_definition_changed"] = True
    if static_input_change_locations:
        details["unqualified_runtime_function_static_input_changed"] = True
        details["unqualified_runtime_function_static_input_change_count"] = len(
            static_input_change_locations
        )
    change = Change(
        "unqualified_runtime_functions_changed",
        None,
        "high",
        details=details,
    )
    finding = Finding(
        "FF075",
        "high",
        (
            "An unqualified runtime-function candidate, relevant formula-defined "
            "name, or statically visible input changed; a formula may now bind to "
            "a different VBA, COM/Automation, XLL, or registered runtime."
        ),
        details=details,
    )
    return [change], [finding]


def _worksheet_code_resource_registration_changes(
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
    static_input_change_locations: set[CellKey],
) -> tuple[list[Change], list[Finding]]:
    """Flag stored ``REGISTER.ID`` calls and their static inputs.

    FormulaFence records only the stored registration expression and relevant
    named-definition chain. It does not resolve a module path, load a DLL/XLL,
    evaluate a formula, or inspect a host's security configuration. Ordinary
    cell edits are in scope only when the static dependency graph reaches a
    registration formula in either snapshot.
    """
    old_registrations: WorksheetCodeResourceRegistrationSnapshot = (
        before.worksheet_code_resource_registrations
    )
    new_registrations: WorksheetCodeResourceRegistrationSnapshot = (
        after.worksheet_code_resource_registrations
    )
    if old_registrations == new_registrations and not static_input_change_locations:
        return [], []

    details: dict[str, object] = {
        "before": old_registrations.to_dict(),
        "after": new_registrations.to_dict(),
    }
    if old_registrations.call_signature != new_registrations.call_signature:
        details["worksheet_code_resource_registration_formula_material_changed"] = True
    if old_registrations.definition_signature != new_registrations.definition_signature:
        details["worksheet_code_resource_registration_definition_material_changed"] = (
            True
        )
    if static_input_change_locations:
        details["worksheet_code_resource_registration_static_input_changed"] = True
        details["worksheet_code_resource_registration_static_input_change_count"] = len(
            static_input_change_locations
        )
    change = Change(
        "worksheet_code_resource_registrations_changed",
        None,
        "high",
        details=details,
    )
    finding = Finding(
        "FF067",
        "high",
        (
            "A worksheet code-resource registration or a statically visible input "
            "changed; Excel may now register a different DLL or code resource."
        ),
        details=details,
    )
    return [change], [finding]


def _formula_defined_xlm_registration_changes(
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
    static_input_change_locations: set[CellKey],
) -> tuple[list[Change], list[Finding]]:
    """Flag XLM ``REGISTER`` calls stored in definitions and their inputs.

    FormulaFence records only the stored formula-defined expression and its
    relevant named-definition chain. It does not resolve a module path, load a
    DLL/XLL, evaluate a formula, execute a macro, or inspect host security
    configuration. Ordinary cell edits are in scope only when the static
    dependency graph reaches a cell that invokes a stored registration in
    either snapshot.
    """
    old_registrations: FormulaDefinedXlmRegistrationSnapshot = (
        before.formula_defined_xlm_registrations
    )
    new_registrations: FormulaDefinedXlmRegistrationSnapshot = (
        after.formula_defined_xlm_registrations
    )
    if old_registrations == new_registrations and not static_input_change_locations:
        return [], []

    details: dict[str, object] = {
        "before": old_registrations.to_dict(),
        "after": new_registrations.to_dict(),
    }
    if old_registrations.invocation_signature != new_registrations.invocation_signature:
        details["formula_defined_xlm_registration_invocation_material_changed"] = True
    if old_registrations.definition_signature != new_registrations.definition_signature:
        details["formula_defined_xlm_registration_definition_material_changed"] = True
    if static_input_change_locations:
        details["formula_defined_xlm_registration_static_input_changed"] = True
        details["formula_defined_xlm_registration_static_input_change_count"] = len(
            static_input_change_locations
        )
    change = Change(
        "formula_defined_xlm_registrations_changed",
        None,
        "high",
        details=details,
    )
    finding = Finding(
        "FF068",
        "high",
        (
            "A formula-defined XLM REGISTER call or a statically visible input "
            "changed; Excel may now register a different DLL function or command, "
            "or load an XLL."
        ),
        details=details,
    )
    return [change], [finding]


def _formula_defined_xlm_evaluation_changes(
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
    static_input_change_locations: set[CellKey],
) -> tuple[list[Change], list[Finding]]:
    """Flag stored XLM `EVALUATE` expressions and visible input changes.

    FormulaFence compares the stored formula-defined expression and its
    relevant named-definition chain privately. It records an ordinary cell edit
    only when the existing static dependency graph reaches an invoking
    formula. It does not evaluate text, parse a runtime-generated expression,
    execute a macro, or infer dependencies hidden inside that expression.
    """
    old_evaluations: FormulaDefinedXlmEvaluationSnapshot = (
        before.formula_defined_xlm_evaluations
    )
    new_evaluations: FormulaDefinedXlmEvaluationSnapshot = (
        after.formula_defined_xlm_evaluations
    )
    if old_evaluations == new_evaluations and not static_input_change_locations:
        return [], []

    details: dict[str, object] = {
        "before": old_evaluations.to_dict(),
        "after": new_evaluations.to_dict(),
    }
    if old_evaluations.invocation_signature != new_evaluations.invocation_signature:
        details["formula_defined_xlm_evaluation_invocation_material_changed"] = True
    if old_evaluations.definition_signature != new_evaluations.definition_signature:
        details["formula_defined_xlm_evaluation_definition_material_changed"] = True
    if static_input_change_locations:
        details["formula_defined_xlm_evaluation_static_input_changed"] = True
        details["formula_defined_xlm_evaluation_static_input_change_count"] = len(
            static_input_change_locations
        )
    change = Change(
        "formula_defined_xlm_evaluations_changed",
        None,
        "high",
        details=details,
    )
    finding = Finding(
        "FF069",
        "high",
        (
            "A formula-defined XLM EVALUATE call or a statically visible input "
            "changed; Excel may now calculate a different expression."
        ),
        details=details,
    )
    return [change], [finding]


def _formula_defined_xlm_action_changes(
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
    static_input_change_locations: set[CellKey],
) -> tuple[list[Change], list[Finding]]:
    """Flag selected stored XLM action calls and visible input changes.

    FormulaFence compares stored action-call spelling and the relevant
    named-definition chain privately. It records an ordinary cell edit only
    when the existing static dependency graph reaches an invoking formula. It
    does not evaluate a formula, resolve an action target or event handler, or
    run a macro, program, DLL entry point, or DDE command.
    """
    old_actions: FormulaDefinedXlmActionSnapshot = before.formula_defined_xlm_actions
    new_actions: FormulaDefinedXlmActionSnapshot = after.formula_defined_xlm_actions
    if old_actions == new_actions and not static_input_change_locations:
        return [], []

    details: dict[str, object] = {
        "before": old_actions.to_dict(),
        "after": new_actions.to_dict(),
    }
    if old_actions.invocation_signature != new_actions.invocation_signature:
        details["formula_defined_xlm_action_invocation_material_changed"] = True
    if old_actions.definition_signature != new_actions.definition_signature:
        details["formula_defined_xlm_action_definition_material_changed"] = True
    if static_input_change_locations:
        details["formula_defined_xlm_action_static_input_changed"] = True
        details["formula_defined_xlm_action_static_input_change_count"] = len(
            static_input_change_locations
        )
    change = Change(
        "formula_defined_xlm_actions_changed",
        None,
        "high",
        details=details,
    )
    finding = Finding(
        "FF073",
        "high",
        (
            "A formula-defined XLM action or event-dispatch call or a statically "
            "visible input changed; Excel may now dispatch a different macro, "
            "program, DLL entry point, DDE command, or event handler."
        ),
        details=details,
    )
    return [change], [finding]


def _formula_defined_xlm_get_cell_changes(
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
    static_input_change_locations: set[CellKey],
) -> tuple[list[Change], list[Finding]]:
    """Flag stored XLM `GET.CELL` calls and visible input changes.

    FormulaFence compares the stored formula-defined information call and its
    relevant named-definition chain privately. It records an ordinary cell edit
    only when the existing static dependency graph reaches an invoking formula.
    It does not evaluate the call, determine its information type, resolve a
    dynamic reference, or simulate Excel display and formatting state.
    """
    old_get_cell_calls: FormulaDefinedXlmGetCellSnapshot = (
        before.formula_defined_xlm_get_cell_calls
    )
    new_get_cell_calls: FormulaDefinedXlmGetCellSnapshot = (
        after.formula_defined_xlm_get_cell_calls
    )
    if old_get_cell_calls == new_get_cell_calls and not static_input_change_locations:
        return [], []

    details: dict[str, object] = {
        "before": old_get_cell_calls.to_dict(),
        "after": new_get_cell_calls.to_dict(),
    }
    if old_get_cell_calls.invocation_signature != new_get_cell_calls.invocation_signature:
        details["formula_defined_xlm_get_cell_invocation_material_changed"] = True
    if old_get_cell_calls.definition_signature != new_get_cell_calls.definition_signature:
        details["formula_defined_xlm_get_cell_definition_material_changed"] = True
    if static_input_change_locations:
        details["formula_defined_xlm_get_cell_static_input_changed"] = True
        details["formula_defined_xlm_get_cell_static_input_change_count"] = len(
            static_input_change_locations
        )
    change = Change(
        "formula_defined_xlm_get_cell_calls_changed",
        None,
        "high",
        details=details,
    )
    finding = Finding(
        "FF070",
        "high",
        (
            "A formula-defined XLM GET.CELL call or a statically visible input "
            "changed; Excel may now calculate from different cell information."
        ),
        details=details,
    )
    return [change], [finding]


def _formula_defined_xlm_environment_information_changes(
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
    static_input_change_locations: set[CellKey],
) -> tuple[list[Change], list[Finding]]:
    """Flag stored XLM environment-information calls and visible inputs.

    FormulaFence compares the stored formula-defined call and its relevant
    named-definition chain privately. It records an ordinary cell edit only
    when the existing static dependency graph reaches an invoking formula. It
    does not evaluate the call, determine its information type, resolve a
    dynamic reference, or simulate workbook, workspace, or document state.
    """
    old_calls: FormulaDefinedXlmEnvironmentInformationSnapshot = (
        before.formula_defined_xlm_environment_information_calls
    )
    new_calls: FormulaDefinedXlmEnvironmentInformationSnapshot = (
        after.formula_defined_xlm_environment_information_calls
    )
    if old_calls == new_calls and not static_input_change_locations:
        return [], []

    details: dict[str, object] = {
        "before": old_calls.to_dict(),
        "after": new_calls.to_dict(),
    }
    if old_calls.invocation_signature != new_calls.invocation_signature:
        details[
            "formula_defined_xlm_environment_information_invocation_material_changed"
        ] = True
    if old_calls.definition_signature != new_calls.definition_signature:
        details[
            "formula_defined_xlm_environment_information_definition_material_changed"
        ] = True
    if static_input_change_locations:
        details[
            "formula_defined_xlm_environment_information_static_input_changed"
        ] = True
        details[
            "formula_defined_xlm_environment_information_static_input_change_count"
        ] = len(static_input_change_locations)
    change = Change(
        "formula_defined_xlm_environment_information_calls_changed",
        None,
        "high",
        details=details,
    )
    finding = Finding(
        "FF071",
        "high",
        (
            "A formula-defined XLM environment-information call or a statically "
            "visible input changed; Excel may now calculate from different "
            "workbook, workspace, or document information."
        ),
        details=details,
    )
    return [change], [finding]


def _formula_environment_information_changes(
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
    static_input_change_locations: set[CellKey],
) -> tuple[list[Change], list[Finding]]:
    """Flag native workbook/environment formula material and visible inputs.

    FormulaFence compares stored calls and their relevant named-definition
    chains privately. It records an ordinary cell edit only when the existing
    static dependency graph reaches an invoking formula. When a complete raw
    workbook tab catalog changes, it also records stored SHEET/SHEETS calls
    that might now observe different tab information. It does not evaluate a
    call, determine an information type, resolve a dynamic argument, or
    simulate file, client, workspace, selection, or workbook state.
    """
    old_calls: FormulaEnvironmentInformationSnapshot = (
        before.formula_environment_information_calls
    )
    new_calls: FormulaEnvironmentInformationSnapshot = (
        after.formula_environment_information_calls
    )
    workbook_tab_catalog_changed = (
        before.workbook_tab_order_complete
        and after.workbook_tab_order_complete
        and before.workbook_tab_order != after.workbook_tab_order
    )
    has_sheet_catalog_calls = bool(
        old_calls.sheet_function_count
        or old_calls.implicit_sheets_reference_function_count
        or new_calls.sheet_function_count
        or new_calls.implicit_sheets_reference_function_count
    )
    workbook_structure_information_changed = (
        workbook_tab_catalog_changed and has_sheet_catalog_calls
    )
    if (
        old_calls == new_calls
        and not static_input_change_locations
        and not workbook_structure_information_changed
    ):
        return [], []

    details: dict[str, object] = {
        "before": old_calls.to_dict(),
        "after": new_calls.to_dict(),
    }
    if old_calls.invocation_signature != new_calls.invocation_signature:
        details["formula_environment_information_invocation_material_changed"] = True
    if old_calls.definition_signature != new_calls.definition_signature:
        details["formula_environment_information_definition_material_changed"] = True
    if static_input_change_locations:
        details["formula_environment_information_static_input_changed"] = True
        details["formula_environment_information_static_input_change_count"] = len(
            static_input_change_locations
        )
    if workbook_structure_information_changed:
        details["formula_environment_information_workbook_tab_catalog_changed"] = True
    change = Change(
        "formula_environment_information_calls_changed",
        None,
        "high",
        details=details,
    )
    finding = Finding(
        "FF072",
        "high",
        (
            "A native CELL, INFO, SHEET, or SHEETS formula call, a statically "
            "visible input, or the workbook tab catalog changed; Excel may now "
            "calculate from different cell, file, client, or workbook-structure "
            "information."
        ),
        details=details,
    )
    return [change], [finding]


def _rich_text_run_entry_map(
    snapshot: RichTextRunSnapshot,
) -> dict[CellKey, RichTextRunEntry]:
    """Index private rich-string presentation records by their effective cells."""
    return {entry.location: entry for entry in snapshot.entries}


def _rich_text_run_controls_changed(
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
) -> tuple[list[Change], list[Finding]]:
    """Flag character-level presentation changes invisible to normal cell diffs.

    A reader normally exposes the concatenated cell text, not the formatting of
    its individual runs. This check avoids duplicating ordinary text edits:
    changing text inside an unchanged run-property sequence stays a normal cell
    change, while a changed property set or a moved style boundary is guarded.
    """
    old_controls = before.rich_text_runs
    new_controls = after.rich_text_runs
    if old_controls == new_controls:
        return [], []

    old_entries = _rich_text_run_entry_map(old_controls)
    new_entries = _rich_text_run_entry_map(new_controls)
    changed_locations: set[CellKey] = set()
    for location in set(old_entries) | set(new_entries):
        old_entry = old_entries.get(location)
        new_entry = new_entries.get(location)
        if old_entry is None or new_entry is None:
            entry = old_entry or new_entry
            if entry is not None and entry.style_sequence_signature is not None:
                changed_locations.add(location)
            continue
        if old_entry.style_sequence_signature != new_entry.style_sequence_signature:
            changed_locations.add(location)
            continue
        if (
            old_entry.text_signature == new_entry.text_signature
            and old_entry.style_layout_signature != new_entry.style_layout_signature
        ):
            changed_locations.add(location)

    unrecognized_metadata_changed = (
        old_controls.unrecognized_rich_text_count
        != new_controls.unrecognized_rich_text_count
    )
    unrecognized_material_changed = (
        (old_controls.unrecognized_rich_text_count or new_controls.unrecognized_rich_text_count)
        and old_controls.definition_signature != new_controls.definition_signature
    )
    if (
        not changed_locations
        and not unrecognized_metadata_changed
        and not unrecognized_material_changed
    ):
        return [], []

    details: dict[str, object] = {
        "before": old_controls.to_dict(),
        "after": new_controls.to_dict(),
        "rich_text_run_control_change_count": len(changed_locations),
    }
    if old_controls.definition_signature != new_controls.definition_signature:
        details["rich_text_run_definition_material_changed"] = True
    if unrecognized_metadata_changed or unrecognized_material_changed:
        details["unrecognized_rich_text_metadata_changed"] = True
    change = Change(
        "rich_text_run_controls_changed",
        None,
        "high",
        details=details,
    )
    finding = Finding(
        "FF043",
        "high",
        (
            "Rich-text run controls changed; character-level formatting or "
            "phonetic hints may be made less visible or misleading."
        ),
        details=details,
    )
    return [change], [finding]


def _legacy_comment_controls_changed(
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
) -> tuple[list[Change], list[Finding]]:
    """Flag legacy note changes that ordinary cell diffs cannot expose."""
    old_comments: LegacyCommentSnapshot = before.legacy_comments
    new_comments: LegacyCommentSnapshot = after.legacy_comments
    if old_comments == new_comments:
        return [], []

    details: dict[str, object] = {
        "before": old_comments.to_dict(),
        "after": new_comments.to_dict(),
    }
    if old_comments.declaration_signature != new_comments.declaration_signature:
        details["legacy_comment_binding_changed"] = True
    if old_comments.definition_signature != new_comments.definition_signature:
        details["legacy_comment_definition_material_changed"] = True
    if old_comments.note_shape_signature != new_comments.note_shape_signature:
        details["legacy_note_vml_material_changed"] = True
    if old_comments.relationship_signature != new_comments.relationship_signature:
        details["legacy_note_relationships_changed"] = True
    if (
        old_comments.unrecognized_legacy_comment_count
        != new_comments.unrecognized_legacy_comment_count
        or (
            (
                old_comments.unrecognized_legacy_comment_count
                or new_comments.unrecognized_legacy_comment_count
            )
            and (
                old_comments.definition_signature != new_comments.definition_signature
                or old_comments.note_shape_signature
                != new_comments.note_shape_signature
            )
        )
    ):
        details["unrecognized_legacy_comment_metadata_changed"] = True
    change = Change(
        "legacy_comment_controls_changed",
        None,
        "high",
        details=details,
    )
    finding = Finding(
        "FF046",
        "high",
        (
            "Legacy Excel notes or threaded-comment placeholders changed; review "
            "text, author context, cell association, visibility, or layout may "
            "be altered outside cells."
        ),
        details=details,
    )
    return [change], [finding]


def _cell_hyperlink_controls_changed(
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
) -> tuple[list[Change], list[Finding]]:
    """Flag stored cell hyperlink changes that ordinary values cannot expose."""
    old_hyperlinks: CellHyperlinkSnapshot = before.cell_hyperlinks
    new_hyperlinks: CellHyperlinkSnapshot = after.cell_hyperlinks
    if old_hyperlinks == new_hyperlinks:
        return [], []

    details: dict[str, object] = {
        "before": old_hyperlinks.to_dict(),
        "after": new_hyperlinks.to_dict(),
    }
    if old_hyperlinks.declaration_signature != new_hyperlinks.declaration_signature:
        details["cell_hyperlink_binding_changed"] = True
    if old_hyperlinks.definition_signature != new_hyperlinks.definition_signature:
        details["cell_hyperlink_definition_material_changed"] = True
    if old_hyperlinks.relationship_signature != new_hyperlinks.relationship_signature:
        details["cell_hyperlink_relationships_changed"] = True
    if (
        old_hyperlinks.unrecognized_cell_hyperlink_count
        != new_hyperlinks.unrecognized_cell_hyperlink_count
        or (
            (
                old_hyperlinks.unrecognized_cell_hyperlink_count
                or new_hyperlinks.unrecognized_cell_hyperlink_count
            )
            and (
                old_hyperlinks.definition_signature
                != new_hyperlinks.definition_signature
                or old_hyperlinks.relationship_signature
                != new_hyperlinks.relationship_signature
            )
        )
    ):
        details["unrecognized_cell_hyperlink_metadata_changed"] = True
    change = Change(
        "cell_hyperlink_controls_changed",
        None,
        "high",
        details=details,
    )
    finding = Finding(
        "FF047",
        "high",
        (
            "Worksheet cell hyperlinks changed; a reviewer may be redirected or "
            "shown a different target outside the ordinary cell value."
        ),
        details=details,
    )
    return [change], [finding]


def _worksheet_sparkline_controls_changed(
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
) -> tuple[list[Change], list[Finding]]:
    """Flag stored sparkline changes that ordinary cell values cannot expose."""
    old_sparklines: WorksheetSparklineSnapshot = before.worksheet_sparklines
    new_sparklines: WorksheetSparklineSnapshot = after.worksheet_sparklines
    if old_sparklines == new_sparklines:
        return [], []

    details: dict[str, object] = {
        "before": old_sparklines.to_dict(),
        "after": new_sparklines.to_dict(),
    }
    if old_sparklines.binding_signature != new_sparklines.binding_signature:
        details["worksheet_sparkline_bindings_changed"] = True
    if old_sparklines.definition_signature != new_sparklines.definition_signature:
        details["worksheet_sparkline_definition_material_changed"] = True
    if (
        old_sparklines.unrecognized_worksheet_sparkline_count
        != new_sparklines.unrecognized_worksheet_sparkline_count
        or (
            (
                old_sparklines.unrecognized_worksheet_sparkline_count
                or new_sparklines.unrecognized_worksheet_sparkline_count
            )
            and old_sparklines.definition_signature
            != new_sparklines.definition_signature
        )
    ):
        details["unrecognized_worksheet_sparkline_metadata_changed"] = True
    change = Change(
        "worksheet_sparkline_controls_changed",
        None,
        "high",
        details=details,
    )
    finding = Finding(
        "FF048",
        "high",
        (
            "Worksheet sparklines changed; a compact visual trend summary may "
            "now show different source data, output cells, or rendering controls."
        ),
        details=details,
    )
    return [change], [finding]


def _xml_mapping_controls_changed(
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
) -> tuple[list[Change], list[Finding]]:
    """Flag XML-map changes that can redirect import/export data outside cells."""
    old_mappings: XmlMappingSnapshot = before.xml_mapping_controls
    new_mappings: XmlMappingSnapshot = after.xml_mapping_controls
    if old_mappings == new_mappings:
        return [], []

    details: dict[str, object] = {
        "before": old_mappings.to_dict(),
        "after": new_mappings.to_dict(),
    }
    if old_mappings.declaration_signature != new_mappings.declaration_signature:
        details["xml_mapping_declarations_changed"] = True
    if old_mappings.binding_signature != new_mappings.binding_signature:
        details["xml_mapping_bindings_changed"] = True
    if old_mappings.relationship_signature != new_mappings.relationship_signature:
        details["xml_mapping_relationships_changed"] = True
    if (
        old_mappings.unrecognized_xml_mapping_count
        != new_mappings.unrecognized_xml_mapping_count
        or (
            (
                old_mappings.unrecognized_xml_mapping_count
                or new_mappings.unrecognized_xml_mapping_count
            )
            and (
                old_mappings.declaration_signature
                != new_mappings.declaration_signature
                or old_mappings.binding_signature
                != new_mappings.binding_signature
                or old_mappings.relationship_signature
                != new_mappings.relationship_signature
            )
        )
    ):
        details["unrecognized_xml_mapping_metadata_changed"] = True
    change = Change(
        "xml_mapping_controls_changed",
        None,
        "high",
        details=details,
    )
    finding = Finding(
        "FF049",
        "high",
        (
            "XML-mapped workbook controls changed; a refresh or export may now "
            "route operational data through a different schema, binding, or target."
        ),
        details=details,
    )
    return [change], [finding]


def _digital_signature_controls_changed(
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
) -> tuple[list[Change], list[Finding]]:
    """Flag package- and VBA-signature envelope changes outside normal cells."""
    old_signatures: DigitalSignatureSnapshot = before.digital_signatures
    new_signatures: DigitalSignatureSnapshot = after.digital_signatures
    if old_signatures == new_signatures:
        return [], []

    details: dict[str, object] = {
        "before": old_signatures.to_dict(),
        "after": new_signatures.to_dict(),
    }
    if (
        old_signatures.package_signature_signature
        != new_signatures.package_signature_signature
    ):
        details["package_signature_material_changed"] = True
    if (
        old_signatures.package_signature_coverage
        != new_signatures.package_signature_coverage
    ):
        details["package_signature_manifest_coverage_changed"] = True
    if (
        old_signatures.vba_signature_payload_signature
        != new_signatures.vba_signature_payload_signature
    ):
        details["vba_project_signature_payload_changed"] = True
    if old_signatures.relationship_signature != new_signatures.relationship_signature:
        details["digital_signature_relationships_changed"] = True
    if (
        old_signatures.unrecognized_digital_signature_count
        != new_signatures.unrecognized_digital_signature_count
        or (
            (
                old_signatures.unrecognized_digital_signature_count
                or new_signatures.unrecognized_digital_signature_count
            )
            and (
                old_signatures.package_signature_signature
                != new_signatures.package_signature_signature
                or old_signatures.vba_signature_payload_signature
                != new_signatures.vba_signature_payload_signature
                or old_signatures.relationship_signature
                != new_signatures.relationship_signature
            )
        )
    ):
        details["unrecognized_digital_signature_metadata_changed"] = True
    change = Change(
        "digital_signature_controls_changed",
        None,
        "high",
        details=details,
    )
    finding = Finding(
        "FF050",
        "high",
        (
            "Digital-signature controls changed; package or VBA provenance and "
            "integrity-assurance metadata may have been added, removed, or altered."
        ),
        details=details,
    )
    return [change], [finding]


def _rich_data_controls_changed(
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
) -> tuple[list[Change], list[Finding]]:
    """Flag rich entity data and related relationship changes outside cells."""
    old_rich_data: RichDataSnapshot = before.rich_data
    new_rich_data: RichDataSnapshot = after.rich_data
    if old_rich_data == new_rich_data:
        return [], []

    details: dict[str, object] = {
        "before": old_rich_data.to_dict(),
        "after": new_rich_data.to_dict(),
    }
    if old_rich_data.definition_signature != new_rich_data.definition_signature:
        details["rich_data_definitions_changed"] = True
    if old_rich_data.value_signature != new_rich_data.value_signature:
        details["rich_data_values_changed"] = True
    if (
        old_rich_data.metadata_binding_signature
        != new_rich_data.metadata_binding_signature
    ):
        details["rich_data_metadata_bindings_changed"] = True
    if old_rich_data.relationship_signature != new_rich_data.relationship_signature:
        details["rich_data_relationships_changed"] = True
    if (
        old_rich_data.unrecognized_rich_data_count
        != new_rich_data.unrecognized_rich_data_count
        or (
            (
                old_rich_data.unrecognized_rich_data_count
                or new_rich_data.unrecognized_rich_data_count
            )
            and (
                old_rich_data.definition_signature
                != new_rich_data.definition_signature
                or old_rich_data.value_signature != new_rich_data.value_signature
                or (
                    old_rich_data.metadata_binding_signature
                    != new_rich_data.metadata_binding_signature
                )
                or (
                    old_rich_data.relationship_signature
                    != new_rich_data.relationship_signature
                )
            )
        )
    ):
        details["unrecognized_rich_data_metadata_changed"] = True
    change = Change(
        "rich_data_controls_changed",
        None,
        "high",
        details=details,
    )
    finding = Finding(
        "FF051",
        "high",
        (
            "Rich data controls changed; workbook entity values, provider-linked "
            "data, external image associations, or metadata bindings may have "
            "been added, removed, or altered."
        ),
        details=details,
    )
    return [change], [finding]


def _custom_data_store_controls_changed(
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
) -> tuple[list[Change], list[Finding]]:
    """Flag persisted custom workbook state that ordinary cells do not expose."""
    old_stores: CustomDataStoreSnapshot = before.custom_data_stores
    new_stores: CustomDataStoreSnapshot = after.custom_data_stores
    if old_stores == new_stores:
        return [], []

    details: dict[str, object] = {
        "before": old_stores.to_dict(),
        "after": new_stores.to_dict(),
    }
    if old_stores.custom_xml_signature != new_stores.custom_xml_signature:
        details["custom_xml_state_changed"] = True
    if old_stores.custom_data_signature != new_stores.custom_data_signature:
        details["custom_data_material_changed"] = True
    if (
        old_stores.document_property_signature
        != new_stores.document_property_signature
    ):
        details["document_custom_properties_changed"] = True
    if old_stores.relationship_signature != new_stores.relationship_signature:
        details["custom_data_store_relationships_changed"] = True
    if (
        old_stores.unrecognized_custom_data_store_count
        != new_stores.unrecognized_custom_data_store_count
        or (
            (
                old_stores.unrecognized_custom_data_store_count
                or new_stores.unrecognized_custom_data_store_count
            )
            and (
                old_stores.custom_xml_signature != new_stores.custom_xml_signature
                or old_stores.custom_data_signature != new_stores.custom_data_signature
                or (
                    old_stores.document_property_signature
                    != new_stores.document_property_signature
                )
                or (
                    old_stores.relationship_signature
                    != new_stores.relationship_signature
                )
            )
        )
    ):
        details["unrecognized_custom_data_store_metadata_changed"] = True
    change = Change(
        "custom_data_store_changed",
        None,
        "high",
        details=details,
    )
    finding = Finding(
        "FF052",
        "high",
        (
            "Custom workbook data stores changed; persisted add-in state, "
            "custom binary data, or custom document properties may have been "
            "added, removed, or altered outside ordinary cells."
        ),
        details=details,
    )
    return [change], [finding]


def _threaded_comment_controls_changed(
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
) -> tuple[list[Change], list[Finding]]:
    """Flag modern-comment changes that ordinary cell diffs cannot expose."""
    old_comments: ThreadedCommentSnapshot = before.threaded_comments
    new_comments: ThreadedCommentSnapshot = after.threaded_comments
    if old_comments == new_comments:
        return [], []

    details: dict[str, object] = {
        "before": old_comments.to_dict(),
        "after": new_comments.to_dict(),
    }
    if old_comments.declaration_signature != new_comments.declaration_signature:
        details["threaded_comment_binding_changed"] = True
    if old_comments.definition_signature != new_comments.definition_signature:
        details["threaded_comment_definition_material_changed"] = True
    if old_comments.person_signature != new_comments.person_signature:
        details["threaded_comment_person_material_changed"] = True
    if old_comments.relationship_signature != new_comments.relationship_signature:
        details["threaded_comment_relationships_changed"] = True
    if (
        old_comments.unrecognized_threaded_comment_count
        != new_comments.unrecognized_threaded_comment_count
        or (
            (
                old_comments.unrecognized_threaded_comment_count
                or new_comments.unrecognized_threaded_comment_count
            )
            and (
                old_comments.definition_signature != new_comments.definition_signature
                or old_comments.person_signature != new_comments.person_signature
            )
        )
    ):
        details["unrecognized_threaded_comment_metadata_changed"] = True
    change = Change(
        "threaded_comment_controls_changed",
        None,
        "high",
        details=details,
    )
    finding = Finding(
        "FF045",
        "high",
        (
            "Modern threaded comments changed; discussion text, replies, resolution, "
            "mentions, or collaborator bindings may be altered outside cells."
        ),
        details=details,
    )
    return [change], [finding]


def _worksheet_drawing_shape_controls_changed(
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
) -> tuple[list[Change], list[Finding]]:
    """Flag DrawingML shape, connector, or graphic-frame changes outside cells."""
    old_shapes: WorksheetDrawingShapeSnapshot = before.worksheet_drawing_shapes
    new_shapes: WorksheetDrawingShapeSnapshot = after.worksheet_drawing_shapes
    if old_shapes == new_shapes:
        return [], []

    details: dict[str, object] = {
        "before": old_shapes.to_dict(),
        "after": new_shapes.to_dict(),
    }
    if old_shapes.declaration_signature != new_shapes.declaration_signature:
        details["worksheet_drawing_shape_binding_changed"] = True
    if old_shapes.definition_signature != new_shapes.definition_signature:
        details["worksheet_drawing_shape_definition_material_changed"] = True
    if old_shapes.diagram_signature != new_shapes.diagram_signature:
        details["worksheet_drawing_diagram_material_changed"] = True
    if old_shapes.relationship_signature != new_shapes.relationship_signature:
        details["worksheet_drawing_shape_relationships_changed"] = True
    if (
        old_shapes.unrecognized_shape_count != new_shapes.unrecognized_shape_count
        or (
            (old_shapes.unrecognized_shape_count or new_shapes.unrecognized_shape_count)
            and old_shapes.definition_signature != new_shapes.definition_signature
        )
    ):
        details["unrecognized_worksheet_drawing_shape_metadata_changed"] = True
    change = Change(
        "worksheet_drawing_shape_controls_changed",
        None,
        "high",
        details=details,
    )
    finding = Finding(
        "FF044",
        "high",
        (
            "Worksheet DrawingML shape, connector, or non-chart graphic-frame controls "
            "changed; text, visual cues, anchors, diagrams, attachments, or linked "
            "actions may be altered outside cells."
        ),
        details=details,
    )
    return [change], [finding]


def _worksheet_image_controls_changed(
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
) -> tuple[list[Change], list[Finding]]:
    """Flag native worksheet image changes outside ordinary cells."""
    old_images: WorksheetImageSnapshot = before.worksheet_images
    new_images: WorksheetImageSnapshot = after.worksheet_images
    if old_images == new_images:
        return [], []

    details: dict[str, object] = {
        "before": old_images.to_dict(),
        "after": new_images.to_dict(),
    }
    if old_images.declaration_signature != new_images.declaration_signature:
        details["worksheet_image_binding_changed"] = True
    if old_images.definition_signature != new_images.definition_signature:
        details["worksheet_image_definition_material_changed"] = True
    if old_images.relationship_signature != new_images.relationship_signature:
        details["worksheet_image_relationships_changed"] = True
    if old_images.image_payload_signature != new_images.image_payload_signature:
        details["worksheet_image_payload_material_changed"] = True
    if (
        old_images.unrecognized_image_count != new_images.unrecognized_image_count
        or (
            old_images.unrecognized_image_count or new_images.unrecognized_image_count
        )
        and (
            old_images.definition_signature != new_images.definition_signature
            or old_images.image_payload_signature != new_images.image_payload_signature
        )
    ):
        details["unrecognized_worksheet_image_metadata_changed"] = True
    change = Change(
        "worksheet_image_controls_changed",
        None,
        "high",
        details=details,
    )
    finding = Finding(
        "FF059",
        "high",
        (
            "Native worksheet images changed; floating pictures, sheet backgrounds, "
            "or header/footer watermarks may alter a report outside cells."
        ),
        details=details,
    )
    return [change], [finding]


def compare_snapshots(
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
    *,
    max_change_analysis_states: int = DEFAULT_MAX_CHANGE_ANALYSIS_STATES,
    _state_budget: ChangeAnalysisBudget | None = None,
) -> DiffReport:
    """Compare workbook semantics and attach bounded local impact to each edit."""
    before.require_full_inspection("Workbook comparison")
    after.require_full_inspection("Workbook comparison")
    if max_change_analysis_states < 1:
        raise FormulaFenceError("max_change_analysis_states must be at least 1.")
    state_budget = _state_budget or ChangeAnalysisBudget(
        max_states=max_change_analysis_states,
        scope="Workbook",
    )
    changes: list[Change] = []
    findings: list[Finding] = []
    formula_changed_locations: set[CellKey] = set()
    semantic_cell_changes: list[Change] = []
    formula_external_action_static_input_changes: set[CellKey] = set()
    formula_dde_link_static_input_changes: set[CellKey] = set()
    python_in_excel_static_input_changes: set[CellKey] = set()
    office_custom_function_static_input_changes: set[CellKey] = set()
    unqualified_runtime_function_static_input_changes: set[CellKey] = set()
    worksheet_code_resource_registration_static_input_changes: set[CellKey] = set()
    formula_defined_xlm_registration_static_input_changes: set[CellKey] = set()
    formula_defined_xlm_evaluation_static_input_changes: set[CellKey] = set()
    formula_defined_xlm_action_static_input_changes: set[CellKey] = set()
    formula_defined_xlm_get_cell_static_input_changes: set[CellKey] = set()
    formula_defined_xlm_environment_information_static_input_changes: set[
        CellKey
    ] = set()
    formula_environment_information_static_input_changes: set[CellKey] = set()
    formula_external_action_cells = (
        before.formula_external_actions.action_cells
        | after.formula_external_actions.action_cells
    )
    formula_dde_link_cells = (
        before.formula_dde_links.dde_cells | after.formula_dde_links.dde_cells
    )
    python_in_excel_cells = (
        before.python_in_excel.python_cells | after.python_in_excel.python_cells
    )
    office_custom_function_cells = (
        before.office_custom_functions.call_cells
        | after.office_custom_functions.call_cells
    )
    unqualified_runtime_function_cells = (
        before.unqualified_runtime_functions.call_cells
        | after.unqualified_runtime_functions.call_cells
    )
    worksheet_code_resource_registration_cells = (
        before.worksheet_code_resource_registrations.registration_cells
        | after.worksheet_code_resource_registrations.registration_cells
    )
    formula_defined_xlm_registration_cells = (
        before.formula_defined_xlm_registrations.registration_cells
        | after.formula_defined_xlm_registrations.registration_cells
    )
    formula_defined_xlm_evaluation_cells = (
        before.formula_defined_xlm_evaluations.evaluation_cells
        | after.formula_defined_xlm_evaluations.evaluation_cells
    )
    formula_defined_xlm_action_cells = (
        before.formula_defined_xlm_actions.action_cells
        | after.formula_defined_xlm_actions.action_cells
    )
    formula_defined_xlm_get_cell_cells = (
        before.formula_defined_xlm_get_cell_calls.get_cell_cells
        | after.formula_defined_xlm_get_cell_calls.get_cell_cells
    )
    formula_defined_xlm_environment_information_cells = (
        before.formula_defined_xlm_environment_information_calls.environment_information_cells
        | after.formula_defined_xlm_environment_information_calls.environment_information_cells
    )
    formula_environment_information_cells = (
        before.formula_environment_information_calls.environment_information_cells
        | after.formula_environment_information_calls.environment_information_cells
    )

    all_locations = sorted(set(before.cells) | set(after.cells), key=_location_sort_key)
    for location in all_locations:
        old_cell = before.cells.get(location)
        new_cell = after.cells.get(location)
        classified = _cell_change_kind(old_cell, new_cell)
        if classified is None:
            continue
        kind, severity = classified
        impact_analysis = analyze_downstream_impact(
            location,
            before,
            after,
            state_budget=state_budget,
        )
        impact = impact_analysis.impacted
        sampled_impacts = tuple(sorted(impact, key=_location_sort_key)[:_IMPACT_SAMPLE_SIZE])
        details: dict[str, object] = {}
        if sampled_impacts:
            details["impact_paths"] = _serialise_impact_paths(
                sampled_impacts, impact_analysis.paths
            )
        if impact_analysis.truncated:
            details["impact_truncated_at"] = _IMPACT_NODE_LIMIT
        change = Change(
            kind=kind,
            location=location,
            severity=severity,
            before=old_cell,
            after=new_cell,
            impact_count=len(impact),
            impacted_cells=sampled_impacts,
            details=details,
        )
        changes.append(change)
        semantic_cell_changes.append(change)
        if formula_external_action_cells & impact:
            formula_external_action_static_input_changes.add(location)
        if formula_dde_link_cells & impact:
            formula_dde_link_static_input_changes.add(location)
        if python_in_excel_cells & impact:
            python_in_excel_static_input_changes.add(location)
        if office_custom_function_cells & impact:
            office_custom_function_static_input_changes.add(location)
        if unqualified_runtime_function_cells & impact:
            unqualified_runtime_function_static_input_changes.add(location)
        if worksheet_code_resource_registration_cells & impact:
            worksheet_code_resource_registration_static_input_changes.add(location)
        if formula_defined_xlm_registration_cells & impact:
            formula_defined_xlm_registration_static_input_changes.add(location)
        if formula_defined_xlm_evaluation_cells & impact:
            formula_defined_xlm_evaluation_static_input_changes.add(location)
        if formula_defined_xlm_action_cells & impact:
            formula_defined_xlm_action_static_input_changes.add(location)
        if formula_defined_xlm_get_cell_cells & impact:
            formula_defined_xlm_get_cell_static_input_changes.add(location)
        if formula_defined_xlm_environment_information_cells & impact:
            formula_defined_xlm_environment_information_static_input_changes.add(
                location
            )
        if formula_environment_information_cells & impact:
            formula_environment_information_static_input_changes.add(location)
        if new_cell is not None and new_cell.is_formula:
            formula_changed_locations.add(location)
        if old_cell is not None and old_cell.is_formula:
            formula_changed_locations.add(location)
        if kind == "formula_to_value":
            findings.append(
                Finding(
                    "FF001",
                    "high",
                    "Formula was replaced with a value.",
                    location,
                    details={
                        "impact_count": len(impact),
                        "impact_paths": details.get("impact_paths", []),
                    },
                )
            )
        elif kind == "formula_removed":
            findings.append(
                Finding(
                    "FF002",
                    "high",
                    "Formula cell was removed or blanked.",
                    location,
                    details={
                        "impact_count": len(impact),
                        "impact_paths": details.get("impact_paths", []),
                    },
                )
            )

    formula_cached_result_changes, formula_cached_result_findings = (
        _formula_cached_result_changes(
            before,
            after,
            semantic_cell_changes,
            state_budget=state_budget,
        )
    )
    changes.extend(formula_cached_result_changes)
    findings.extend(formula_cached_result_findings)

    formula_external_action_changes, formula_external_action_findings = (
        _formula_external_action_changes(
            before,
            after,
            formula_external_action_static_input_changes,
        )
    )
    changes.extend(formula_external_action_changes)
    findings.extend(formula_external_action_findings)

    formula_dde_link_changes, formula_dde_link_findings = _formula_dde_link_changes(
        before,
        after,
        formula_dde_link_static_input_changes,
    )
    changes.extend(formula_dde_link_changes)
    findings.extend(formula_dde_link_findings)

    python_in_excel_changes, python_in_excel_findings = _python_in_excel_changes(
        before,
        after,
        python_in_excel_static_input_changes,
    )
    changes.extend(python_in_excel_changes)
    findings.extend(python_in_excel_findings)

    office_custom_function_changes, office_custom_function_findings = (
        _office_custom_function_changes(
            before,
            after,
            office_custom_function_static_input_changes,
        )
    )
    changes.extend(office_custom_function_changes)
    findings.extend(office_custom_function_findings)

    (
        unqualified_runtime_function_changes,
        unqualified_runtime_function_findings,
    ) = _unqualified_runtime_function_changes(
        before,
        after,
        unqualified_runtime_function_static_input_changes,
    )
    changes.extend(unqualified_runtime_function_changes)
    findings.extend(unqualified_runtime_function_findings)

    (
        worksheet_code_resource_registration_changes,
        worksheet_code_resource_registration_findings,
    ) = _worksheet_code_resource_registration_changes(
        before,
        after,
        worksheet_code_resource_registration_static_input_changes,
    )
    changes.extend(worksheet_code_resource_registration_changes)
    findings.extend(worksheet_code_resource_registration_findings)

    (
        formula_defined_xlm_registration_changes,
        formula_defined_xlm_registration_findings,
    ) = _formula_defined_xlm_registration_changes(
        before,
        after,
        formula_defined_xlm_registration_static_input_changes,
    )
    changes.extend(formula_defined_xlm_registration_changes)
    findings.extend(formula_defined_xlm_registration_findings)

    (
        formula_defined_xlm_evaluation_changes,
        formula_defined_xlm_evaluation_findings,
    ) = _formula_defined_xlm_evaluation_changes(
        before,
        after,
        formula_defined_xlm_evaluation_static_input_changes,
    )
    changes.extend(formula_defined_xlm_evaluation_changes)
    findings.extend(formula_defined_xlm_evaluation_findings)

    (
        formula_defined_xlm_action_changes,
        formula_defined_xlm_action_findings,
    ) = _formula_defined_xlm_action_changes(
        before,
        after,
        formula_defined_xlm_action_static_input_changes,
    )
    changes.extend(formula_defined_xlm_action_changes)
    findings.extend(formula_defined_xlm_action_findings)

    (
        formula_defined_xlm_get_cell_changes,
        formula_defined_xlm_get_cell_findings,
    ) = _formula_defined_xlm_get_cell_changes(
        before,
        after,
        formula_defined_xlm_get_cell_static_input_changes,
    )
    changes.extend(formula_defined_xlm_get_cell_changes)
    findings.extend(formula_defined_xlm_get_cell_findings)

    (
        formula_defined_xlm_environment_information_changes,
        formula_defined_xlm_environment_information_findings,
    ) = _formula_defined_xlm_environment_information_changes(
        before,
        after,
        formula_defined_xlm_environment_information_static_input_changes,
    )
    changes.extend(formula_defined_xlm_environment_information_changes)
    findings.extend(formula_defined_xlm_environment_information_findings)

    (
        formula_environment_information_changes,
        formula_environment_information_findings,
    ) = _formula_environment_information_changes(
        before,
        after,
        formula_environment_information_static_input_changes,
    )
    changes.extend(formula_environment_information_changes)
    findings.extend(formula_environment_information_findings)

    rich_text_run_changes, rich_text_run_findings = _rich_text_run_controls_changed(
        before,
        after,
    )
    changes.extend(rich_text_run_changes)
    findings.extend(rich_text_run_findings)

    cell_hyperlink_changes, cell_hyperlink_findings = _cell_hyperlink_controls_changed(
        before,
        after,
    )
    changes.extend(cell_hyperlink_changes)
    findings.extend(cell_hyperlink_findings)

    worksheet_sparkline_changes, worksheet_sparkline_findings = (
        _worksheet_sparkline_controls_changed(before, after)
    )
    changes.extend(worksheet_sparkline_changes)
    findings.extend(worksheet_sparkline_findings)

    xml_mapping_changes, xml_mapping_findings = _xml_mapping_controls_changed(
        before,
        after,
    )
    changes.extend(xml_mapping_changes)
    findings.extend(xml_mapping_findings)

    digital_signature_changes, digital_signature_findings = (
        _digital_signature_controls_changed(before, after)
    )
    changes.extend(digital_signature_changes)
    findings.extend(digital_signature_findings)

    rich_data_changes, rich_data_findings = _rich_data_controls_changed(before, after)
    changes.extend(rich_data_changes)
    findings.extend(rich_data_findings)

    custom_data_store_changes, custom_data_store_findings = (
        _custom_data_store_controls_changed(before, after)
    )
    changes.extend(custom_data_store_changes)
    findings.extend(custom_data_store_findings)

    legacy_comment_changes, legacy_comment_findings = _legacy_comment_controls_changed(
        before,
        after,
    )
    changes.extend(legacy_comment_changes)
    findings.extend(legacy_comment_findings)

    threaded_comment_changes, threaded_comment_findings = _threaded_comment_controls_changed(
        before,
        after,
    )
    changes.extend(threaded_comment_changes)
    findings.extend(threaded_comment_findings)

    worksheet_drawing_shape_changes, worksheet_drawing_shape_findings = (
        _worksheet_drawing_shape_controls_changed(before, after)
    )
    changes.extend(worksheet_drawing_shape_changes)
    findings.extend(worksheet_drawing_shape_findings)

    worksheet_image_changes, worksheet_image_findings = _worksheet_image_controls_changed(
        before,
        after,
    )
    changes.extend(worksheet_image_changes)
    findings.extend(worksheet_image_findings)

    if before.external_workbook_link_surfaces != after.external_workbook_link_surfaces:
        old_link_surfaces: ExternalWorkbookLinkSurfaceSnapshot = (
            before.external_workbook_link_surfaces
        )
        new_link_surfaces: ExternalWorkbookLinkSurfaceSnapshot = (
            after.external_workbook_link_surfaces
        )
        details: dict[str, object] = {
            "before": old_link_surfaces.to_dict(),
            "after": new_link_surfaces.to_dict(),
        }
        if old_link_surfaces.ledger_signature != new_link_surfaces.ledger_signature:
            details["external_workbook_link_surface_material_changed"] = True
        if (
            old_link_surfaces.opaque_chart_part_count
            != new_link_surfaces.opaque_chart_part_count
            or (
                old_link_surfaces.opaque_chart_part_count
                or new_link_surfaces.opaque_chart_part_count
            )
            and old_link_surfaces.ledger_signature != new_link_surfaces.ledger_signature
        ):
            details["opaque_chart_link_surface_coverage_changed"] = True
        changes.append(
            Change(
                "external_workbook_link_surfaces_changed",
                None,
                "high",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF081",
                "high",
                (
                    "Static external-workbook link surfaces changed; worksheet formulas, "
                    "defined names, data validation, or chart formulas may now bind a "
                    "different source or target."
                ),
                details=details,
            )
        )

    newly_broken = after.broken_references - before.broken_references
    for location in sorted(newly_broken, key=_location_sort_key):
        findings.append(
            Finding("FF003", "critical", "Formula contains a broken #REF! reference.", location)
        )
    new_external = after.external_references - before.external_references
    for location in sorted(new_external, key=_location_sort_key):
        findings.append(
            Finding(
                "FF004",
                "high",
                "Formula introduces an explicit external-workbook reference.",
                location,
            )
        )

    for location, tokens in _new_coverage_items(
        before.unresolved_reference_tokens, after.unresolved_reference_tokens
    ):
        details = {"tokens": list(tokens)}
        changes.append(
            Change(
                "unresolved_formula_reference_added",
                location,
                "medium",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF011",
                "medium",
                "Formula introduces a reference that FormulaFence cannot statically resolve.",
                location,
                details=details,
            )
        )

    for location, functions in _new_coverage_items(
        before.dynamic_reference_functions, after.dynamic_reference_functions
    ):
        details = {"functions": list(functions)}
        changes.append(
            Change(
                "dynamic_formula_reference_added",
                location,
                "medium",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF012",
                "medium",
                (
                    "Formula introduces a dynamic reference function; "
                    "dependency impact may be incomplete."
                ),
                location,
                details=details,
            )
        )

    for location, tokens in _new_coverage_items(
        before.spill_reference_tokens, after.spill_reference_tokens
    ):
        details = {"tokens": list(tokens)}
        changes.append(
            Change(
                "spill_reference_added",
                location,
                "medium",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF015",
                "medium",
                (
                    "Formula introduces a dynamic-array spill reference; the anchor is "
                    "traced but dynamic extent and blockers may affect impact."
                ),
                location,
                details=details,
            )
        )

    for location in sorted(
        after.tokenization_failure_cells - before.tokenization_failure_cells,
        key=_location_sort_key,
    ):
        changes.append(
            Change(
                "formula_tokenization_failure_added",
                location,
                "medium",
            )
        )
        findings.append(
            Finding(
                "FF016",
                "medium",
                "Formula could not be tokenized; dependency impact may be incomplete.",
                location,
            )
        )

    for location, tokens in _new_coverage_items(
        before.implicit_intersection_tokens, after.implicit_intersection_tokens
    ):
        details = {"tokens": list(tokens)}
        changes.append(
            Change(
                "implicit_intersection_added",
                location,
                "medium",
                details=details,
            )
        )
        findings.append(
            Finding(
                "FF017",
                "medium",
                (
                    "Formula introduces explicit implicit intersection; it can change which "
                    "cell a range or array contributes."
                ),
                location,
                details=details,
            )
        )

    array_formula_changes, array_formula_findings = _array_formula_semantics_changes(
        before,
        after,
        state_budget=state_budget,
    )
    changes.extend(array_formula_changes)
    findings.extend(array_formula_findings)
    dynamic_array_reference_changes, dynamic_array_reference_findings = (
        _dynamic_array_output_reference_changes(
            before,
            after,
            state_budget=state_budget,
        )
    )
    changes.extend(dynamic_array_reference_changes)
    findings.extend(dynamic_array_reference_findings)
    control_changes, control_findings = _workbook_control_changes(before, after)
    changes.extend(control_changes)
    findings.extend(control_findings)
    findings.extend(_formula_pattern_findings(after, formula_changed_locations))

    changes.sort(key=lambda change: (_location_sort_key(change.location), change.kind))
    findings.sort(key=lambda finding: (_location_sort_key(finding.location), finding.rule_id))
    return DiffReport(
        before=before,
        after=after,
        changes=changes,
        findings=findings,
        formula_external_action_static_input_cells=frozenset(
            formula_external_action_static_input_changes
        ),
        formula_dde_link_static_input_cells=frozenset(
            formula_dde_link_static_input_changes
        ),
        python_in_excel_static_input_cells=frozenset(
            python_in_excel_static_input_changes
        ),
        office_custom_function_static_input_cells=frozenset(
            office_custom_function_static_input_changes
        ),
        unqualified_runtime_function_static_input_cells=frozenset(
            unqualified_runtime_function_static_input_changes
        ),
        worksheet_code_resource_registration_static_input_cells=frozenset(
            worksheet_code_resource_registration_static_input_changes
        ),
        formula_defined_xlm_registration_static_input_cells=frozenset(
            formula_defined_xlm_registration_static_input_changes
        ),
        formula_defined_xlm_evaluation_static_input_cells=frozenset(
            formula_defined_xlm_evaluation_static_input_changes
        ),
        formula_defined_xlm_action_static_input_cells=frozenset(
            formula_defined_xlm_action_static_input_changes
        ),
        formula_defined_xlm_get_cell_static_input_cells=frozenset(
            formula_defined_xlm_get_cell_static_input_changes
        ),
        formula_defined_xlm_environment_information_static_input_cells=frozenset(
            formula_defined_xlm_environment_information_static_input_changes
        ),
        formula_environment_information_static_input_cells=frozenset(
            formula_environment_information_static_input_changes
        ),
    )


def report_severities(report: DiffReport, extra_findings: Iterable[Finding] = ()) -> list[str]:
    """Return both risk-finding and semantic-change severities for CLI thresholds."""
    return [
        *(change.severity for change in report.changes),
        *(finding.severity for finding in report.findings),
        *(finding.severity for finding in extra_findings),
    ]
